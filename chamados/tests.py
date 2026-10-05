"""Testes do app Chamados (pytest + pytest-django).

Cobre os testes críticos da ARCHITECTURE.md: máquina de estados (transições
válidas/inválidas), imutabilidade dos fatos de abertura, campos obrigatórios por
ação, reabertura derivada do log, fronteira da Inteligência, permissão de
abertura, protocolo por ano, painel derivado do log e atomicidade.
"""
import pytest
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied, ValidationError
from django.urls import reverse

from acompanhamento.models import Clientes
from produto.models import Produto
from chamados import services
from chamados.enums import Acao, Setor, Status
from chamados.models import Chamado, ChamadoEquipamento, ChamadoEvento
from chamados.selectors import metricas_painel

User = get_user_model()


# --------------------------------------------------------------------------- #
# Fixtures                                                                     #
# --------------------------------------------------------------------------- #


@pytest.fixture
def grupos(db):
    quality, _ = Group.objects.get_or_create(name="quality")
    inteligencia, _ = Group.objects.get_or_create(name="inteligencia")
    return quality, inteligencia


@pytest.fixture
def user_quality(db, grupos):
    quality, _ = grupos
    u = User.objects.create_user(username="q1", password="x")
    u.groups.add(quality)
    return u


@pytest.fixture
def outro_quality(db, grupos):
    quality, _ = grupos
    u = User.objects.create_user(username="q2", password="x")
    u.groups.add(quality)
    return u


@pytest.fixture
def user_inteligencia(db, grupos):
    _, inteligencia = grupos
    u = User.objects.create_user(username="i1", password="x")
    u.groups.add(inteligencia)
    return u


@pytest.fixture
def outro_inteligencia(db, grupos):
    _, inteligencia = grupos
    u = User.objects.create_user(username="i2", password="x")
    u.groups.add(inteligencia)
    return u


@pytest.fixture
def user_expedicao(db):
    """Usuário do grupo `expedicao` (fila compartilhada). O grupo é criado por
    migration no app real; aqui garantimos sua existência no banco de teste."""
    expedicao, _ = Group.objects.get_or_create(name="expedicao")
    u = User.objects.create_user(username="e1", password="x")
    u.groups.add(expedicao)
    return u


@pytest.fixture
def outro_expedicao(db):
    expedicao, _ = Group.objects.get_or_create(name="expedicao")
    u = User.objects.create_user(username="e2", password="x")
    u.groups.add(expedicao)
    return u


@pytest.fixture
def user_laboratorio(db):
    """Usuário do grupo `laboratorio` (fila compartilhada que recebe as chegadas)."""
    laboratorio, _ = Group.objects.get_or_create(name="laboratorio")
    u = User.objects.create_user(username="l1", password="x")
    u.groups.add(laboratorio)
    return u


@pytest.fixture
def user_comercial(db):
    """Usuário do grupo `COMERCIAL` (maiúsculo — reaproveitado do sistema)."""
    comercial, _ = Group.objects.get_or_create(name="COMERCIAL")
    u = User.objects.create_user(username="c1", password="x")
    u.groups.add(comercial)
    return u


@pytest.fixture
def user_financeiro(db):
    """Usuário do grupo `financeiro` (fila compartilhada que fatura e encerra)."""
    financeiro, _ = Group.objects.get_or_create(name="financeiro")
    u = User.objects.create_user(username="f1", password="x")
    u.groups.add(financeiro)
    return u


@pytest.fixture
def user_recepcao(db):
    return _usuario_do_grupo("rec1", "recepcao")


@pytest.fixture
def user_configuracao(db):
    return _usuario_do_grupo("cfg1", "CONFIGURACAO")


def _usuario_do_grupo(username, grupo):
    u, _ = User.objects.get_or_create(username=username)
    u.groups.add(Group.objects.get_or_create(name=grupo)[0])
    return u


@pytest.fixture
def user_comum(db):
    return User.objects.create_user(username="comum", password="x")


@pytest.fixture
def cliente(db):
    """Cliente do cadastro (acompanhamento.Clientes) usado nas aberturas."""
    return Clientes.objects.create(nome="ACME", endereco="Rua 1", cnpj="00000000000000")


@pytest.fixture
def produto(db):
    """Modelo de equipamento (produto.Produto) usado nas aberturas."""
    return Produto.objects.create(nome="Rastreador GT06")


@pytest.fixture
def manutencao(db, cliente):
    """Entrada de manutenção elegível para vínculo (mesmos critérios da tela
    "Registro das entradas": status em andamento e data a partir de 2026)."""
    import datetime

    from django.utils import timezone

    from registrodemanutencao.models import registrodemanutencao

    m = registrodemanutencao.objects.create(nome=cliente, status="Manutenção")
    # data_criacao costuma ser auto_now_add; garantimos o corte de data do filtro.
    registrodemanutencao.objects.filter(pk=m.pk).update(
        data_criacao=timezone.make_aware(datetime.datetime(2026, 6, 1, 12, 0))
    )
    m.refresh_from_db()
    return m


def _pdf_falso(nome="termo.pdf"):
    """Arquivo PDF mínimo em memória, para os testes de anexo do termo."""
    from django.core.files.uploadedfile import SimpleUploadedFile

    return SimpleUploadedFile(nome, b"%PDF-1.4 teste", content_type="application/pdf")


def _abrir(autor, responsavel, cliente=None, modelo=None, equipamentos_override=None, **extra):
    if cliente is None:
        cliente = Clientes.objects.create(
            nome="ACME", endereco="Rua 1", cnpj="00000000000000"
        )
    if modelo is None:
        modelo = Produto.objects.create(nome="Rastreador GT06")
    return services.abrir_chamado(
        autor=autor,
        cliente=cliente,
        categoria="HARDWARE",
        equipamentos=equipamentos_override or [("EQ-001", modelo)],
        problema_relatado="Não liga",
        responsavel=responsavel,
        contato_nome="João da Silva",
        contato_telefone="11999998888",
        contato_email="joao@exemplo.com",
        contato_meio="WHATSAPP",
        **extra,
    )


def _blocos(*grupos):
    """POST da abertura: um bloco por (modelo, [números][, customização, contrato]),
    como a tela envia. Sem customização/contrato, usa valores válidos."""
    dados = {"grupo": [str(i) for i in range(len(grupos))]}
    for i, (modelo, numeros, *extra) in enumerate(grupos):
        padrao = ["Termo branco", "Retornavel"] if modelo else ["", ""]  # bloco vazio: selects em branco
        customizacao, contrato = (extra + padrao)[:2]
        dados[f"grupo_{i}_modelo"] = modelo
        dados[f"grupo_{i}_numero"] = numeros
        dados[f"grupo_{i}_customizacao"] = customizacao
        dados[f"grupo_{i}_tipo_contrato"] = contrato
    return dados


# --------------------------------------------------------------------------- #
# Abertura + protocolo                                                         #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_abertura_cria_chamado_aberto_com_evento_inicial(user_quality):
    chamado = _abrir(user_quality, user_quality)
    assert chamado.status == Status.ABERTO
    assert chamado.aberto_por == user_quality
    assert chamado.aberto_em is not None
    evento = chamado.eventos.get()
    assert evento.acao == Acao.ABRIR
    assert evento.estado_origem is None
    assert evento.estado_destino == Status.ABERTO


@pytest.mark.django_db
def test_abertura_por_nao_quality_e_rejeitada(user_inteligencia, user_quality):
    """RN-01 — usuário fora do grupo quality não abre chamado."""
    with pytest.raises(PermissionDenied):
        _abrir(user_inteligencia, user_quality)


@pytest.mark.django_db
def test_responsavel_deve_ser_quality(user_quality, user_inteligencia):
    """RN-02 — o responsável do chamado deve ser do grupo quality."""
    with pytest.raises(ValidationError):
        _abrir(user_quality, user_inteligencia)


@pytest.mark.django_db
def test_abrir_encaminhado_exige_procedimento_tratativa_e_resp_inteligencia(
    user_quality, user_inteligencia
):
    """RN-08 — abrir já ENCAMINHADO exige os três campos."""
    with pytest.raises(ValidationError):
        _abrir(user_quality, user_quality, encaminhar=True)  # sem os campos

    chamado = _abrir(
        user_quality,
        user_quality,
        encaminhar=True,
        procedimento_realizado="tentei reset",
        tratativa="trocar antena",
        responsavel_inteligencia=user_inteligencia,
    )
    assert chamado.status == Status.ENCAMINHADO
    assert chamado.responsavel_inteligencia == user_inteligencia


@pytest.mark.django_db
def test_protocolo_formato_e_sequencial_por_ano(user_quality):
    """RN-07 — AAAA-NNNNNN; sequencial incrementa dentro do ano."""
    c1 = _abrir(user_quality, user_quality)
    c2 = _abrir(user_quality, user_quality)
    ano = c1.aberto_em.year
    assert c1.protocolo == f"{ano}-000001"
    assert c2.protocolo == f"{ano}-000002"


@pytest.mark.django_db
def test_protocolo_reinicia_no_novo_ano(user_quality, cliente, produto):
    """RN-07 — o sequencial reinicia na virada de ano (protocolos de 2024 não
    contam para 2025)."""
    Chamado.objects.create(
        protocolo="2024-000009",
        cliente=cliente, categoria="OUTROS", numero_equipamento="E",
        problema_relatado="p", responsavel=user_quality,
        contato_nome="Fulano", contato_meio="TELEFONE",
        aberto_por=user_quality, aberto_em=_dt(2024),
    )
    assert services.gerar_protocolo(2025) == "2025-000001"


def _dt(ano):
    from django.utils import timezone
    import datetime

    return timezone.make_aware(datetime.datetime(ano, 6, 1, 12, 0))


# --------------------------------------------------------------------------- #
# Máquina de estados — transições válidas e inválidas (RN-14)                  #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_finalizar_exige_procedimento(user_quality):
    """RN-10 — Finalizar sem procedimento falha; com procedimento, resolve."""
    chamado = _abrir(user_quality, user_quality)
    with pytest.raises(ValidationError):
        services.executar(chamado, Acao.FINALIZAR, {}, user_quality)

    services.executar(
        chamado, Acao.FINALIZAR, {"procedimento_realizado": "trocado"}, user_quality
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO


@pytest.mark.django_db
def test_encaminhar_exige_procedimento_tratativa_e_resp(user_quality, user_inteligencia):
    """RN-09 — Encaminhar exige os três campos."""
    chamado = _abrir(user_quality, user_quality)
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.ENCAMINHAR, {"procedimento_realizado": "x"}, user_quality
        )
    services.executar(
        chamado,
        Acao.ENCAMINHAR,
        {
            "procedimento_realizado": "x",
            "tratativa": "y",
            "responsavel_inteligencia": user_inteligencia,
        },
        user_quality,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.ENCAMINHADO
    assert chamado.responsavel_inteligencia == user_inteligencia


@pytest.mark.django_db
def test_resolvido_e_terminal(user_quality):
    """RN-13 — RESOLVIDO não tem transição de saída."""
    chamado = _abrir(user_quality, user_quality)
    services.executar(
        chamado, Acao.FINALIZAR, {"procedimento_realizado": "ok"}, user_quality
    )
    chamado.refresh_from_db()
    for acao in (Acao.ENCAMINHAR, Acao.BLOQUEAR, Acao.REABRIR):
        with pytest.raises(ValidationError):
            services.executar(chamado, acao, {"motivo": "m", "procedimento_realizado": "p", "tratativa": "t"}, user_quality)


@pytest.mark.django_db
def test_transicao_invalida_aberto_para_resolver_intel(user_quality, user_inteligencia):
    """ENCAMINHAR→RESOLVER é da Inteligência; ABERTO não sai para RESOLVER via
    RESOLVER (só via FINALIZAR do Quality)."""
    chamado = _abrir(user_quality, user_quality)
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.RESOLVER, {"procedimento_realizado": "x"}, user_inteligencia
        )


# --------------------------------------------------------------------------- #
# Fronteira da Inteligência (RN-16)                                            #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_inteligencia_resolve_encaminhado(user_quality, user_inteligencia):
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    services.aceitar_tratativa(chamado, user_inteligencia)  # aceite antes de agir
    services.executar(
        chamado, Acao.RESOLVER, {"procedimento_realizado": "resolvido"}, user_inteligencia
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO


@pytest.mark.django_db
def test_inteligencia_nao_pode_agir_em_aberto(user_quality, user_inteligencia):
    """RN-16 — a Inteligência não age em chamado ABERTO (posse do Quality);
    tentar finalizar um ABERTO não é dela."""
    chamado = _abrir(user_quality, user_quality)
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.FINALIZAR, {"procedimento_realizado": "x"}, user_inteligencia
        )


@pytest.mark.django_db
def test_quality_nao_age_apos_encaminhar(user_quality, user_inteligencia):
    """RN-17 — após ENCAMINHADO, a posse é da Inteligência; Quality não resolve."""
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.RESOLVER, {"procedimento_realizado": "x"}, user_quality
        )


# --------------------------------------------------------------------------- #
# Bloqueio e reabertura (RN-11, RN-12) — destino derivado do log (ADR-005)     #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_bloquear_exige_motivo(user_quality):
    chamado = _abrir(user_quality, user_quality)
    with pytest.raises(ValidationError):
        services.executar(chamado, Acao.BLOQUEAR, {}, user_quality)


@pytest.mark.django_db
def test_reabrir_de_aberto_volta_para_aberto(user_quality):
    """RN-12 — bloqueado a partir de ABERTO reabre para ABERTO (dono Quality)."""
    chamado = _abrir(user_quality, user_quality)
    services.executar(chamado, Acao.BLOQUEAR, {"motivo": "peça"}, user_quality)
    chamado.refresh_from_db()
    assert chamado.status == Status.BLOQUEADO
    services.executar(chamado, Acao.REABRIR, {"motivo": "chegou"}, user_quality)
    chamado.refresh_from_db()
    assert chamado.status == Status.ABERTO


@pytest.mark.django_db
def test_reabrir_de_encaminhado_volta_para_encaminhado_posse_intel(
    user_quality, user_inteligencia
):
    """RN-12 — bloqueado a partir de ENCAMINHADO reabre para ENCAMINHADO, e é a
    Inteligência (dono atual) quem reabre."""
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    services.executar(chamado, Acao.BLOQUEAR, {"motivo": "terceiro"}, user_inteligencia)
    chamado.refresh_from_db()
    # Quality não reabre um bloqueio que era da Inteligência.
    with pytest.raises(PermissionDenied):
        services.executar(chamado, Acao.REABRIR, {"motivo": "x"}, user_quality)
    services.executar(chamado, Acao.REABRIR, {"motivo": "resolvido terceiro"}, user_inteligencia)
    chamado.refresh_from_db()
    assert chamado.status == Status.ENCAMINHADO


# --------------------------------------------------------------------------- #
# Imutabilidade dos fatos de abertura (RN-03)                                  #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_transicao_nao_altera_fatos_de_abertura(user_quality):
    """RN-03 — nenhuma ação toca cliente/categoria/equipamento/responsável."""
    chamado = _abrir(user_quality, user_quality)
    antes = (chamado.cliente, chamado.categoria, chamado.numero_equipamento,
             list(chamado.equipamentos.values_list("numero", "modelo_id")),
             chamado.problema_relatado,
             chamado.responsavel_id, chamado.aberto_em,
             chamado.contato_nome, chamado.contato_telefone,
             chamado.contato_email, chamado.contato_meio)
    services.executar(chamado, Acao.BLOQUEAR, {"motivo": "aguardando"}, user_quality)
    chamado.refresh_from_db()
    depois = (chamado.cliente, chamado.categoria, chamado.numero_equipamento,
              list(chamado.equipamentos.values_list("numero", "modelo_id")),
             chamado.problema_relatado,
              chamado.responsavel_id, chamado.aberto_em,
              chamado.contato_nome, chamado.contato_telefone,
              chamado.contato_email, chamado.contato_meio)
    assert antes == depois


@pytest.mark.django_db
def test_evento_e_append_only(user_quality):
    """ADR-010 — um ChamadoEvento já criado não pode ser reescrito nem apagado."""
    chamado = _abrir(user_quality, user_quality)
    evento = chamado.eventos.get()
    evento.motivo = "hack"
    with pytest.raises(ValidationError):
        evento.save()
    with pytest.raises(ValidationError):
        evento.delete()


# --------------------------------------------------------------------------- #
# Painel derivado do log (RN-05, ADR-011)                                      #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_painel_bate_com_o_log(user_quality, user_inteligencia):
    aberto = _abrir(user_quality, user_quality)  # ABERTO
    encaminhado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    resolvido = _abrir(user_quality, user_quality)
    services.executar(resolvido, Acao.FINALIZAR, {"procedimento_realizado": "ok"}, user_quality)

    m = metricas_painel()
    assert m["fila_ativa"] == 1
    assert m["encaminhados"] == 1
    assert m["resolvidos_hoje"] == 1
    assert "em_analise" not in m


# --------------------------------------------------------------------------- #
# Views / permissões (integração)                                             #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_fila_exige_login(client):
    resp = client.get(reverse("chamados:fila"))
    assert resp.status_code == 302
    assert "login" in resp.url.lower() or resp.url


@pytest.mark.django_db
def test_abrir_get_bloqueado_para_nao_quality(client, user_inteligencia):
    """Inteligência é operador, mas ABRIR é exclusivo de quality → 403."""
    client.force_login(user_inteligencia)
    resp = client.get(reverse("chamados:abrir"))
    assert resp.status_code == 403


@pytest.mark.django_db
def test_fila_bloqueada_para_nao_operador(client, user_comum):
    """Usuário logado fora de quality/inteligencia não acessa a fila → 403."""
    client.force_login(user_comum)
    resp = client.get(reverse("chamados:fila"))
    assert resp.status_code == 403


@pytest.mark.django_db
def test_fila_liberada_para_inteligencia(client, user_inteligencia):
    """Inteligência é operador → vê a fila (RN-18)."""
    client.force_login(user_inteligencia)
    resp = client.get(reverse("chamados:fila"))
    assert resp.status_code == 200


@pytest.mark.django_db
def test_detalhe_bloqueado_para_nao_operador(client, user_comum, user_quality):
    chamado = _abrir(user_quality, user_quality)
    client.force_login(user_comum)
    resp = client.get(reverse("chamados:detalhe", args=[chamado.pk]))
    assert resp.status_code == 403


@pytest.mark.django_db
def test_acao_bloqueada_para_nao_operador(client, user_comum, user_quality):
    """Mesmo o POST de ação é barrado na URL para não-operador → 403."""
    chamado = _abrir(user_quality, user_quality)
    client.force_login(user_comum)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.FINALIZAR])
    )
    assert resp.status_code == 403


@pytest.mark.django_db
def test_abrir_post_cria_chamado(client, user_quality, cliente, produto):
    client.force_login(user_quality)
    # `responsavel` não é mais enviado no POST: vem do usuário logado.
    resp = client.post(
        reverse("chamados:abrir"),
        {
            "cliente": cliente.pk,
            "categoria": "HARDWARE",
            **_blocos((produto.pk, ["EQ-9"])),
            "problema_relatado": "Falha",
            "contato_nome": "Maria Contato",
            "contato_telefone": "1133334444",
            "contato_email": "maria@exemplo.com",
            "contato_meio": "EMAIL",
        },
    )
    assert resp.status_code == 302
    chamado = Chamado.objects.get(cliente=cliente, numero_equipamento="EQ-9")
    assert chamado.responsavel == user_quality  # definido pelo usuário logado
    assert chamado.contato_nome == "Maria Contato"
    assert chamado.contato_meio == "EMAIL"


@pytest.mark.django_db
def test_abrir_com_equipamentos_de_modelos_diferentes(client, user_quality, cliente):
    """Cada nº é gravado com o modelo do SEU bloco; bloco vazio é descartado."""
    isca_4g = Produto.objects.create(nome="Isca 4G")
    isca_2g = Produto.objects.create(nome="Isca 2G")
    client.force_login(user_quality)
    resp = client.post(
        reverse("chamados:abrir"),
        {
            "cliente": cliente.pk,
            "categoria": "HARDWARE",
            # o test client envia a lista como múltiplos valores de mesmo name
            **_blocos(
                (isca_4g.pk, ["EQ-1", " EQ-3 ", ""]),
                ("", [""]),  # bloco aberto e deixado vazio: ignorado
                (isca_2g.pk, ["EQ-2"]),
            ),
            "problema_relatado": "Falha",
            "contato_nome": "Contato",
            "contato_meio": "TELEFONE",
        },
    )
    assert resp.status_code == 302
    chamado = Chamado.objects.get(cliente=cliente)
    assert chamado.numero_equipamento == "EQ-1, EQ-3, EQ-2"
    # sabotagem: gravar o 1º modelo em todas as linhas no service → vermelho
    assert list(chamado.equipamentos.values_list("numero", "modelo__nome")) == [
        ("EQ-1", "Isca 4G"), ("EQ-3", "Isca 4G"), ("EQ-2", "Isca 2G"),
    ]


@pytest.mark.django_db
@pytest.mark.parametrize(
    "blocos, erro",
    [
        ([("m", ["EQ-1"]), ("", ["EQ-2", "EQ-3"])], "Selecione o modelo dos equipamentos EQ-2, EQ-3."),
        ([("m", ["EQ-1"]), ("m", [""])], "Informe ao menos um nº para o modelo Rastreador GT06."),
        ([("m", ["EQ-1"]), ("m", ["EQ-1"])], "Nº EQ-1 informado mais de uma vez."),
        ([("", [""])], "Informe ao menos um equipamento."),
        ([("999999", ["EQ-1"])], "Modelo de equipamento inválido."),
    ],
)
def test_abrir_rejeita_linhas_de_equipamento_invalidas(
    client, user_quality, cliente, produto, blocos, erro
):
    """Nada é gravado quando algum bloco modelo + nºs está incompleto ou repete nº."""
    client.force_login(user_quality)
    resp = client.post(
        reverse("chamados:abrir"),
        {
            "cliente": cliente.pk,
            "categoria": "HARDWARE",
            **_blocos(*[(produto.pk if m == "m" else m, nums) for m, nums in blocos]),
            "problema_relatado": "Falha",
            "contato_nome": "Contato",
            "contato_meio": "TELEFONE",
        },
    )
    assert resp.status_code == 200
    assert erro in resp.context["form"].errors["equipamentos"]
    assert not Chamado.objects.exists()


@pytest.mark.django_db
def test_servico_rejeita_numero_repetido_sem_gravar(user_quality, produto):
    """O service reforça a regra do form: nº repetido não cria o chamado."""
    with pytest.raises(ValidationError):
        _abrir(user_quality, user_quality, equipamentos_override=[
            ("EQ-1", produto), ("EQ-1", produto),
        ])
    assert not Chamado.objects.exists()
    assert not ChamadoEquipamento.objects.exists()


@pytest.mark.django_db
def test_abrir_ignora_responsavel_forjado_no_post(
    client, user_quality, outro_quality, cliente, produto
):
    """O responsável é sempre o usuário logado; um `responsavel` enviado no POST
    (tentando forjar outro Quality) é ignorado pela view."""
    client.force_login(user_quality)
    resp = client.post(
        reverse("chamados:abrir"),
        {
            "cliente": cliente.pk,
            "categoria": "HARDWARE",
            **_blocos((produto.pk, ["EQ-FORJADO"])),
            "problema_relatado": "Falha",
            "responsavel": outro_quality.pk,  # tentativa de forjar — deve ser ignorada
            "contato_nome": "Contato",
            "contato_meio": "TELEFONE",
        },
    )
    assert resp.status_code == 302
    chamado = Chamado.objects.get(numero_equipamento="EQ-FORJADO")
    assert chamado.responsavel == user_quality  # e NÃO outro_quality


@pytest.mark.django_db
def test_acao_finalizar_via_view(client, user_quality):
    chamado = _abrir(user_quality, user_quality)
    client.force_login(user_quality)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.FINALIZAR]),
        {"procedimento_realizado": "resolvido no local"},
    )
    assert resp.status_code == 302
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO


# --------------------------------------------------------------------------- #
# Isolamento da Inteligência — só vê e age no que foi encaminhado a ELE        #
# --------------------------------------------------------------------------- #


def _encaminhado_para(user_quality, intel):
    """Chamado em ENCAMINHADO, já ACEITO pela inteligência.

    O aceite é obrigatório antes de agir (marco inicial do SLA), então os helpers
    já o executam para que os testes de fluxo sigam direto ao ponto.
    """
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=intel,
    )
    services.aceitar_tratativa(chamado, intel)
    return chamado


@pytest.mark.django_db
def test_intel_nao_age_em_encaminhado_de_outro_intel(
    user_quality, user_inteligencia, outro_inteligencia
):
    """Posse individual: intel B não resolve chamado encaminhado ao intel A."""
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.RESOLVER, {"procedimento_realizado": "x"}, outro_inteligencia
        )


@pytest.mark.django_db
def test_intel_dono_resolve_o_seu(user_quality, user_inteligencia):
    """O intel a quem foi encaminhado continua resolvendo normalmente."""
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    services.executar(
        chamado, Acao.RESOLVER, {"procedimento_realizado": "ok"}, user_inteligencia
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO


@pytest.mark.django_db
def test_fila_do_intel_mostra_so_os_dele(
    client, user_quality, user_inteligencia, outro_inteligencia
):
    """A fila do intel A não lista chamados de quality nem os do intel B."""
    meu = _encaminhado_para(user_quality, user_inteligencia)
    do_outro = _encaminhado_para(user_quality, outro_inteligencia)
    aberto_quality = _abrir(user_quality, user_quality)  # ABERTO, não é de intel

    client.force_login(user_inteligencia)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert pks == {meu.pk}
    assert do_outro.pk not in pks
    assert aberto_quality.pk not in pks


@pytest.mark.django_db
def test_quality_ve_tudo_na_fila(
    client, user_quality, user_inteligencia, outro_inteligencia
):
    """Quality continua vendo todos os chamados."""
    a = _abrir(user_quality, user_quality)
    b = _encaminhado_para(user_quality, user_inteligencia)
    c = _encaminhado_para(user_quality, outro_inteligencia)

    client.force_login(user_quality)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert {a.pk, b.pk, c.pk} <= pks


@pytest.mark.django_db
def test_intel_nao_abre_detalhe_de_outro(
    client, user_quality, user_inteligencia, outro_inteligencia
):
    """Acesso direto por URL: intel A recebe 404 no detalhe de chamado do B."""
    do_outro = _encaminhado_para(user_quality, outro_inteligencia)
    client.force_login(user_inteligencia)
    resp = client.get(reverse("chamados:detalhe", args=[do_outro.pk]))
    assert resp.status_code == 404


@pytest.mark.django_db
def test_intel_nao_posta_acao_em_chamado_de_outro(
    client, user_quality, user_inteligencia, outro_inteligencia
):
    """POST direto de ação em chamado alheio → 404 (nem chega ao service)."""
    do_outro = _encaminhado_para(user_quality, outro_inteligencia)
    client.force_login(user_inteligencia)
    resp = client.post(
        reverse("chamados:acao", args=[do_outro.pk, Acao.RESOLVER]),
        {"procedimento_realizado": "x"},
    )
    assert resp.status_code == 404
    do_outro.refresh_from_db()
    assert do_outro.status == Status.ENCAMINHADO  # inalterado


# --------------------------------------------------------------------------- #
# Fluxo Expedição — Inteligência encaminha p/ expedição; grupo expedicao age   #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_intel_encaminha_para_expedicao(user_quality, user_inteligencia):
    """A Inteligência (dona do ENCAMINHADO) envia o chamado para EXPEDICAO."""
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    services.executar(
        chamado,
        Acao.ENCAMINHAR_EXPEDICAO,
        {"procedimento_realizado": "verificado", "tratativa": "precisa manutenção"},
        user_inteligencia,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.EXPEDICAO
    # o responsável de inteligência permanece registrado (RN-15 análogo)
    assert chamado.responsavel_inteligencia == user_inteligencia


@pytest.mark.django_db
def test_encaminhar_expedicao_exige_procedimento_e_tratativa(
    user_quality, user_inteligencia
):
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.ENCAMINHAR_EXPEDICAO,
            {"procedimento_realizado": "só isso"}, user_inteligencia,
        )


@pytest.mark.django_db
def test_quality_nao_encaminha_para_expedicao(user_quality, user_inteligencia):
    """Só a Inteligência (dona do ENCAMINHADO) manda p/ expedição — não o Quality."""
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.ENCAMINHAR_EXPEDICAO,
            {"procedimento_realizado": "p", "tratativa": "t"}, user_quality,
        )


def _em_expedicao(user_quality, user_inteligencia, aceitar=True, user_expedicao=None):
    """Chamado levado até EXPEDICAO. Com `aceitar`, a expedição já aceitou."""
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    services.executar(
        chamado, Acao.ENCAMINHAR_EXPEDICAO,
        {"procedimento_realizado": "p", "tratativa": "t"}, user_inteligencia,
    )
    if aceitar and user_expedicao is not None:
        services.aceitar_tratativa(chamado, user_expedicao)
    return chamado


@pytest.mark.django_db
def test_expedicao_so_marca_chegada_nao_resolve(
    user_quality, user_inteligencia, user_expedicao
):
    """A Expedição NÃO resolve nem bloqueia — só marca chegada. As demais ações
    a partir de EXPEDICAO nem existem (transição inválida)."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    for acao, dados in (
        (Acao.RESOLVER, {"procedimento_realizado": "x"}),
        (Acao.BLOQUEAR, {"motivo": "m"}),
        (Acao.ENCAMINHAR_EXPEDICAO, {"procedimento_realizado": "p", "tratativa": "t"}),
    ):
        with pytest.raises(ValidationError):  # não é transição válida de EXPEDICAO
            services.executar(chamado, acao, dados, user_expedicao)
    chamado.refresh_from_db()
    assert chamado.status == Status.EXPEDICAO  # inalterado


@pytest.mark.django_db
def test_expedicao_marca_chegada_vai_para_laboratorio(
    user_quality, user_inteligencia, user_expedicao
):
    """Marcar chegada leva o chamado de EXPEDICAO para LABORATORIO."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO


@pytest.mark.django_db
def test_qualquer_expedicao_marca_chegada_fila_compartilhada(
    user_quality, user_inteligencia, outro_expedicao
):
    """Fila compartilhada: um segundo membro da expedição também marca chegada."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=outro_expedicao)
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, outro_expedicao)
    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO


@pytest.mark.django_db
def test_nao_expedicao_nao_marca_chegada(
    user_quality, user_inteligencia, user_comum
):
    """Quem não é da expedição não marca chegada (posse do grupo expedicao)."""
    chamado = _em_expedicao(user_quality, user_inteligencia)
    with pytest.raises(PermissionDenied):
        services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_comum)


@pytest.mark.django_db
def test_laboratorio_ve_fila_de_laboratorio(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """Após marcar chegada, o chamado (LABORATORIO) aparece na fila do laboratório
    e continua na da expedição, que já passou por ele."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)

    client.force_login(user_laboratorio)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert pks == {chamado.pk}

    client.force_login(user_expedicao)
    resp = client.get(reverse("chamados:fila"))
    pks_exp = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert chamado.pk in pks_exp


@pytest.mark.django_db
def test_marcar_chegada_via_view(client, user_quality, user_inteligencia, user_expedicao):
    """POST da ação muda o status para LABORATORIO e leva a Expedição direto à
    entrada dos equipamentos, preenchida a partir do chamado."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    client.force_login(user_expedicao)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.MARCAR_CHEGADA])
    )
    assert resp.status_code == 302
    assert resp.url == f"{reverse('FormulariosCreateView')}?chamado={chamado.pk}"
    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO


def _expedicao_pode_criar_entrada(user):
    """Na produção, a permissão vem de outro grupo do usuário da expedição."""
    from django.contrib.auth.models import Permission

    user.user_permissions.add(Permission.objects.get(codename="add_registrodemanutencao"))


def _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao):
    """Chamado com isca 4G (EQ-1, EQ-3) e 2G (EQ-2) cuja chegada foi marcada."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    isca_4g, isca_2g = Produto.objects.create(nome="Isca 4G"), Produto.objects.create(nome="Isca 2G")
    chamado.equipamentos.all().delete()
    ChamadoEquipamento.objects.bulk_create(
        ChamadoEquipamento(chamado=chamado, numero=n, modelo=m)
        for n, m in (("EQ-1", isca_4g), ("EQ-2", isca_2g), ("EQ-3", isca_4g))
    )
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    return chamado, isca_4g, isca_2g


def _post_entrada(chamado, *blocos):
    dados = {"nome": chamado.cliente_id, "tipo_entrada": "Manutenção", "status": "Pendente",
             "entregue_por_retirado_por": "", "observacoes": "", "chamado": chamado.pk,
             "item": [str(i) for i in range(len(blocos))]}
    for i, (produto, numeros) in enumerate(blocos):
        dados.update({f"item_{i}_tipo_produto": produto.pk, f"item_{i}_numero": numeros})
    return dados


@pytest.mark.django_db
def test_entrada_vem_preenchida_do_chamado(client, user_quality, user_inteligencia, user_expedicao):
    chamado, isca_4g, isca_2g = _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao)
    _expedicao_pode_criar_entrada(user_expedicao)
    client.force_login(user_expedicao)

    form = client.get(f"{reverse('FormulariosCreateView')}?chamado={chamado.pk}").context["form"]

    assert form.initial["nome"] == chamado.cliente_id
    assert form.blocos_itens() == [
        {"tipo_produto": str(isca_4g.pk), "numero": "EQ-1\nEQ-3", "customizacao": "", "tipo_contrato": ""},
        {"tipo_produto": str(isca_2g.pk), "numero": "EQ-2", "customizacao": "", "tipo_contrato": ""},
    ]


@pytest.mark.django_db
def test_salvar_entrada_vincula_ao_chamado(client, user_quality, user_inteligencia, user_expedicao):
    chamado, isca_4g, isca_2g = _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao)
    _expedicao_pode_criar_entrada(user_expedicao)
    client.force_login(user_expedicao)

    resp = client.post(reverse("FormulariosCreateView"),
                       _post_entrada(chamado, (isca_4g, "EQ-1 EQ-3"), (isca_2g, "EQ-2")))

    chamado.refresh_from_db()
    assert resp.url == reverse("chamados:fila")
    assert chamado.manutencao.itens.count() == 2


@pytest.mark.django_db
def test_vinculo_falho_nao_deixa_entrada_orfa(user_quality, user_inteligencia, user_expedicao, manutencao):
    """Chamado vinculado por outra aba no meio do caminho: a entrada nova é desfeita."""
    from registrodemanutencao.models import ItemEntrada, registrodemanutencao
    from registrodemanutencao.services import criar_entrada

    chamado, isca_4g, _ = _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao)
    Chamado.objects.filter(pk=chamado.pk).update(manutencao=manutencao)
    antes = registrodemanutencao.objects.count()

    with pytest.raises(ValidationError):
        criar_entrada(registrodemanutencao(nome=chamado.cliente), [
            ItemEntrada(tipo_produto=isca_4g, numero_equipamento="EQ-1", quantidade=1)
        ], chamado=chamado)
    assert registrodemanutencao.objects.count() == antes


@pytest.mark.django_db
def test_entrada_nao_vem_do_chamado_para_quem_nao_e_expedicao_nem_laboratorio(
    client, user_quality, user_inteligencia, user_expedicao, user_comercial
):
    chamado, *_ = _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao)
    _expedicao_pode_criar_entrada(user_comercial)
    client.force_login(user_comercial)

    resp = client.get(f"{reverse('FormulariosCreateView')}?chamado={chamado.pk}")

    assert resp.context["chamado"] is None
    assert "nome" not in resp.context["form"].initial


@pytest.mark.django_db
def test_intel_nao_age_apos_enviar_para_expedicao(user_quality, user_inteligencia):
    """Após EXPEDICAO a Inteligência não resolve mais: RESOLVER nem é transição
    válida a partir de EXPEDICAO (a expedição só marca chegada)."""
    chamado = _em_expedicao(user_quality, user_inteligencia)
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.RESOLVER,
            {"procedimento_realizado": "x"}, user_inteligencia,
        )


@pytest.mark.django_db
def test_expedicao_ve_fila_de_expedicao(
    client, user_quality, user_inteligencia, user_expedicao
):
    """A expedição vê os chamados em EXPEDICAO e não os ABERTO/ENCAMINHADO."""
    em_expedicao = _encaminhado_para(user_quality, user_inteligencia)
    services.executar(
        em_expedicao, Acao.ENCAMINHAR_EXPEDICAO,
        {"procedimento_realizado": "p", "tratativa": "t"}, user_inteligencia,
    )
    so_aberto = _abrir(user_quality, user_quality)  # ABERTO
    so_encaminhado = _encaminhado_para(user_quality, user_inteligencia)  # ENCAMINHADO

    client.force_login(user_expedicao)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert pks == {em_expedicao.pk}
    assert so_aberto.pk not in pks
    assert so_encaminhado.pk not in pks


@pytest.mark.django_db
def test_encaminhar_para_expedicao_via_view(
    client, user_quality, user_inteligencia
):
    """POST da ação pela view muda o status para EXPEDICAO."""
    chamado = _encaminhado_para(user_quality, user_inteligencia)
    client.force_login(user_inteligencia)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.ENCAMINHAR_EXPEDICAO]),
        {"procedimento_realizado": "verificado", "tratativa": "manutenção"},
    )
    assert resp.status_code == 302
    chamado.refresh_from_db()
    assert chamado.status == Status.EXPEDICAO


# --------------------------------------------------------------------------- #
# Fluxo Laboratório → Comercial — lab dá a tratativa e encaminha ao comercial  #
# --------------------------------------------------------------------------- #


def _em_laboratorio(user_quality, user_inteligencia, user_expedicao,
                    user_laboratorio=None):
    """Chamado levado até LABORATORIO (aceito pelo lab quando informado)."""
    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    if user_laboratorio is not None:
        services.aceitar_tratativa(chamado, user_laboratorio)
    return chamado


def _em_laboratorio_multi(user_quality, user_inteligencia, user_expedicao, numeros,
                          user_laboratorio=None):
    """Chamado com N equipamentos (numero_equipamento juntado por vírgula), levado
    até LABORATORIO (aceito pelo lab quando informado)."""
    cliente = Clientes.objects.create(nome="ACME", endereco="R", cnpj="00000000000000")
    chamado = _abrir(
        user_quality, user_quality, cliente=cliente,
        encaminhar=True, procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    services.aceitar_tratativa(chamado, user_inteligencia)
    # sobrescreve numero_equipamento com a lista desejada
    chamado.numero_equipamento = ", ".join(numeros)
    chamado.save(update_fields=["numero_equipamento"])
    services.executar(chamado, Acao.ENCAMINHAR_EXPEDICAO,
                      {"procedimento_realizado": "p", "tratativa": "t"}, user_inteligencia)
    services.aceitar_tratativa(chamado, user_expedicao)
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    if user_laboratorio is not None:
        services.aceitar_tratativa(chamado, user_laboratorio)
    return chamado


@pytest.mark.django_db
def test_lab_encaminha_para_comercial_por_equipamento(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """O Laboratório informa a tratativa de CADA equipamento; o chamado vai p/
    COMERCIAL e uma linha de TratativaEquipamento é gravada por equipamento."""
    from chamados.models import TratativaEquipamento

    chamado = _em_laboratorio_multi(
        user_quality, user_inteligencia, user_expedicao, ["EQ-1", "EQ-2", "EQ-3"],
        user_laboratorio=user_laboratorio,
    )
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {"tratativas_equipamento": [
            {"numero": "EQ-1", "tratativa": "trocada a placa"},
            {"numero": "EQ-2", "tratativa": "limpeza de contato"},
            {"numero": "EQ-3", "tratativa": "sem reparo"},
        ]},
        user_laboratorio,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.COMERCIAL

    linhas = TratativaEquipamento.objects.filter(chamado=chamado).order_by("id")
    assert [(l.numero_equipamento, l.tratativa) for l in linhas] == [
        ("EQ-1", "trocada a placa"),
        ("EQ-2", "limpeza de contato"),
        ("EQ-3", "sem reparo"),
    ]
    # a tratativa consolidada reúne os três
    assert "EQ-1: trocada a placa" in chamado.tratativa
    assert "EQ-3: sem reparo" in chamado.tratativa


@pytest.mark.django_db
def test_encaminhar_comercial_exige_tratativa_de_cada_equipamento(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """Falta a tratativa de um equipamento → erro (nem transiciona)."""
    from chamados.models import TratativaEquipamento

    chamado = _em_laboratorio_multi(
        user_quality, user_inteligencia, user_expedicao, ["EQ-1", "EQ-2"],
        user_laboratorio=user_laboratorio,
    )
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.ENCAMINHAR_COMERCIAL,
            {"tratativas_equipamento": [{"numero": "EQ-1", "tratativa": "só esse"}]},
            user_laboratorio,
        )
    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO  # inalterado
    assert TratativaEquipamento.objects.filter(chamado=chamado).count() == 0  # atômico


@pytest.mark.django_db
def test_nao_laboratorio_nao_encaminha_para_comercial(
    user_quality, user_inteligencia, user_expedicao, user_comercial
):
    """Quem não é do laboratório não encaminha p/ comercial (posse do grupo lab)."""
    chamado = _em_laboratorio(user_quality, user_inteligencia, user_expedicao)
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.ENCAMINHAR_COMERCIAL,
            {"tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "x"}]},
            user_comercial,
        )


@pytest.mark.django_db
def test_comercial_ve_fila_de_comercial(
    client, user_quality, user_inteligencia, user_expedicao,
    user_laboratorio, user_comercial,
):
    """Após encaminhar p/ comercial, o chamado (COMERCIAL) aparece na fila do
    comercial e continua na do laboratório, que já passou por ele."""
    chamado = _em_laboratorio(user_quality, user_inteligencia, user_expedicao, user_laboratorio=user_laboratorio)
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {"tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}]},
        user_laboratorio,
    )

    client.force_login(user_comercial)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert pks == {chamado.pk}

    client.force_login(user_laboratorio)
    resp = client.get(reverse("chamados:fila"))
    pks_lab = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert chamado.pk in pks_lab


@pytest.mark.django_db
def test_encaminhar_comercial_via_view_por_equipamento(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    manutencao,
):
    """POST pela view: os campos tratativa_<i> viram linhas por equipamento, a
    manutenção é vinculada e o status vai para COMERCIAL (o lab segue no detalhe)."""
    from chamados.models import TratativaEquipamento

    chamado = _em_laboratorio_multi(
        user_quality, user_inteligencia, user_expedicao, ["EQ-1", "EQ-2"],
        user_laboratorio=user_laboratorio,
    )
    client.force_login(user_laboratorio)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.ENCAMINHAR_COMERCIAL]),
        {
            "tratativa_0": "reparo A", "tratativa_1": "reparo B",
            "manutencao": manutencao.pk,
        },
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])
    chamado.refresh_from_db()
    assert chamado.status == Status.COMERCIAL
    assert chamado.manutencao == manutencao  # vínculo gravado
    linhas = TratativaEquipamento.objects.filter(chamado=chamado).order_by("id")
    assert [(l.numero_equipamento, l.tratativa) for l in linhas] == [
        ("EQ-1", "reparo A"), ("EQ-2", "reparo B"),
    ]


@pytest.mark.django_db
def test_encaminhar_comercial_exige_manutencao(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """Sem selecionar a manutenção, o encaminhamento ao comercial não acontece."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    client.force_login(user_laboratorio)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.ENCAMINHAR_COMERCIAL]),
        {"tratativa_0": "reparo"},  # sem `manutencao`
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])  # erro no form
    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO  # inalterado
    assert chamado.manutencao is None


@pytest.mark.django_db
def test_select_de_manutencoes_lista_as_da_tela_de_entradas(manutencao, cliente):
    """O select do modal usa o mesmo critério da tela "Registro das entradas"."""
    import datetime

    from django.utils import timezone

    from chamados.forms import EncaminharComercialForm
    from chamados.selectors import manutencoes_para_vinculo
    from registrodemanutencao.models import registrodemanutencao

    # fora do filtro: status não listado
    fora_status = registrodemanutencao.objects.create(nome=cliente, status="Finalizado")
    registrodemanutencao.objects.filter(pk=fora_status.pk).update(
        data_criacao=timezone.make_aware(datetime.datetime(2026, 6, 1, 12, 0))
    )
    # fora do filtro: anterior a 2026
    fora_data = registrodemanutencao.objects.create(nome=cliente, status="Manutenção")
    registrodemanutencao.objects.filter(pk=fora_data.pk).update(
        data_criacao=timezone.make_aware(datetime.datetime(2025, 12, 31, 12, 0))
    )

    ids = list(manutencoes_para_vinculo().values_list("id", flat=True))
    assert manutencao.pk in ids
    assert fora_status.pk not in ids
    assert fora_data.pk not in ids

    # rótulo do select: "#ID · Empresa"
    form = EncaminharComercialForm(equipamentos=["EQ-1"])
    rotulo = form.fields["manutencao"].label_from_instance(manutencao)
    assert rotulo == f"#{manutencao.pk} · {cliente.nome}"


@pytest.mark.django_db
def test_encaminhar_comercial_via_view_campo_vazio_falha(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """Deixar a tratativa de um equipamento em branco no POST → não transiciona."""
    chamado = _em_laboratorio_multi(
        user_quality, user_inteligencia, user_expedicao, ["EQ-1", "EQ-2"],
        user_laboratorio=user_laboratorio,
    )
    client.force_login(user_laboratorio)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.ENCAMINHAR_COMERCIAL]),
        {"tratativa_0": "reparo A", "tratativa_1": ""},
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])  # volta ao detalhe c/ erro
    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO  # inalterado


def _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio,
                  numeros=None, user_comercial=None):
    """Chamado levado até COMERCIAL (aceito pelo comercial quando informado)."""
    if numeros:
        chamado = _em_laboratorio_multi(
            user_quality, user_inteligencia, user_expedicao, numeros,
            user_laboratorio=user_laboratorio,
        )
        tratativas = [{"numero": n, "tratativa": f"lab {n}"} for n in numeros]
    else:
        chamado = _em_laboratorio(
            user_quality, user_inteligencia, user_expedicao,
            user_laboratorio=user_laboratorio,
        )
        tratativas = [{"numero": "EQ-001", "tratativa": "lab"}]
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {"tratativas_equipamento": tratativas}, user_laboratorio,
    )
    if user_comercial is not None:
        services.aceitar_tratativa(chamado, user_comercial)
    return chamado


@pytest.mark.django_db
def test_comercial_tem_acao_finalizar(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """O comercial vê a ação 'Finalizar chamado' no estado COMERCIAL."""
    from chamados.selectors import acoes_disponiveis

    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial=user_comercial)
    assert acoes_disponiveis(user_comercial, chamado) == [Acao.FINALIZAR_COMERCIAL]


@pytest.mark.django_db
def test_comercial_finaliza_com_tratativa_e_custo(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """Realizar tratativa grava tratativa_comercial + custo por equipamento."""
    from chamados.models import TratativaEquipamento

    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        numeros=["EQ-1", "EQ-2"], user_comercial=user_comercial,
    )
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {
            "finalizacao_equipamento": [
                {"numero": "EQ-1", "tratativa": "orçado", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"},
                {"numero": "EQ-2", "tratativa": "garantia", "custo": "SEM_CUSTO", "destino": "SUBSTITUICAO"},
            ],
            # Há COM_CUSTO → termo obrigatório.
            "termo_substituicao": _pdf_falso(),
        },
        user_comercial,
    )
    chamado.refresh_from_db()
    # Há equipamento de SUBSTITUIÇÃO → segue para a RECEPÇÃO (abre a requisição).
    assert chamado.status == Status.RECEPCAO

    linhas = {l.numero_equipamento: l for l in
              TratativaEquipamento.objects.filter(chamado=chamado)}
    assert linhas["EQ-1"].tratativa_comercial == "orçado"
    assert linhas["EQ-1"].custo == "COM_CUSTO"
    assert linhas["EQ-2"].tratativa_comercial == "garantia"
    assert linhas["EQ-2"].custo == "SEM_CUSTO"
    # não duplicou linhas: continua 1 por equipamento (a do lab foi completada)
    assert TratativaEquipamento.objects.filter(chamado=chamado).count() == 2


@pytest.mark.django_db
def test_finalizar_comercial_exige_custo_de_cada_equipamento(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        numeros=["EQ-1", "EQ-2"], user_comercial=user_comercial,
    )
    with pytest.raises(ValidationError):  # falta custo do EQ-2
        services.executar(
            chamado, Acao.FINALIZAR_COMERCIAL,
            {"finalizacao_equipamento": [
                {"numero": "EQ-1", "tratativa": "ok", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"},
                {"numero": "EQ-2", "tratativa": "ok", "custo": "", "destino": "SUBSTITUICAO"},
            ]},
            user_comercial,
        )
    chamado.refresh_from_db()
    assert chamado.status == Status.COMERCIAL  # inalterado


@pytest.mark.django_db
def test_nao_comercial_nao_finaliza(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comum
):
    """Quem não é do comercial não finaliza (posse do grupo comercial)."""
    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio)
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.FINALIZAR_COMERCIAL,
            {"finalizacao_equipamento": [
                {"numero": "EQ-001", "tratativa": "x", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
            ]},
            user_comum,
        )


@pytest.mark.django_db
def test_finalizar_comercial_via_view(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """POST pela view: campos tratativa_<i>/custo_<i> viram os dados por equipamento
    e o chamado segue para a Recepção."""
    from chamados.models import TratativaEquipamento

    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        numeros=["EQ-1", "EQ-2"], user_comercial=user_comercial,
    )
    client.force_login(user_comercial)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.FINALIZAR_COMERCIAL]),
        {
            "tratativa_0": "reparo A", "custo_0": "COM_CUSTO", "destino_0": "SUBSTITUICAO",
            "tratativa_1": "reparo B", "custo_1": "SEM_CUSTO", "destino_1": "SUBSTITUICAO",
            "termo_substituicao": _pdf_falso(),  # há COM_CUSTO
        },
    )
    assert resp.status_code == 302
    chamado.refresh_from_db()
    # Há SUBSTITUIÇÃO → vai à RECEPÇÃO.
    assert chamado.status == Status.RECEPCAO
    assert chamado.termo_substituicao  # anexado
    linhas = {l.numero_equipamento: l for l in
              TratativaEquipamento.objects.filter(chamado=chamado)}
    assert linhas["EQ-1"].custo == "COM_CUSTO"
    assert linhas["EQ-2"].custo == "SEM_CUSTO"


@pytest.mark.django_db
def test_resolvido_apos_comercial_e_terminal(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """Encerrado pelo Financeiro, RESOLVIDO é terminal (sem novas ações)."""
    from chamados.selectors import acoes_disponiveis

    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial=user_comercial)
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {"finalizacao_equipamento": [
            {"numero": "EQ-001", "tratativa": "t", "custo": "SEM_CUSTO", "destino": "DEVOLUCAO"}
        ]},
        user_comercial,
    )
    _seguir_ate_financeiro(chamado, user_expedicao)
    financeiro = _usuario_do_grupo("fin1", "financeiro")
    services.aceitar_tratativa(chamado, financeiro)
    services.executar(chamado, Acao.CONFIRMAR_ENCERRAMENTO, {}, financeiro)
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO
    assert acoes_disponiveis(financeiro, chamado) == []


# --------------------------------------------------------------------------- #
# Revisão de visibilidade: tela ÚNICA + detalhe bloqueado por papel            #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_expedicao_continua_no_detalhe_apos_marcar_chegada(
    client, user_quality, user_inteligencia, user_expedicao
):
    """Depois de marcar chegada (vai p/ LABORATORIO) a expedição ainda abre o detalhe."""
    chamado = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    client.force_login(user_expedicao)
    assert client.get(reverse("chamados:detalhe", args=[chamado.pk])).status_code == 200

    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    assert client.get(reverse("chamados:detalhe", args=[chamado.pk])).status_code == 200


@pytest.mark.django_db
def test_laboratorio_continua_no_detalhe_apos_encaminhar_comercial(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """Depois de encaminhar p/ comercial o laboratório ainda abre o detalhe."""
    chamado = _em_laboratorio(user_quality, user_inteligencia, user_expedicao, user_laboratorio=user_laboratorio)
    client.force_login(user_laboratorio)
    assert client.get(reverse("chamados:detalhe", args=[chamado.pk])).status_code == 200

    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {"tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}]},
        user_laboratorio,
    )
    assert client.get(reverse("chamados:detalhe", args=[chamado.pk])).status_code == 200


@pytest.mark.django_db
def test_expedicao_nao_abre_detalhe_de_encaminhado(
    client, user_quality, user_inteligencia, user_expedicao
):
    """A expedição não enxerga (404) um chamado que ainda está ENCAMINHADO."""
    encaminhado = _encaminhado_para(user_quality, user_inteligencia)
    client.force_login(user_expedicao)
    resp = client.get(reverse("chamados:detalhe", args=[encaminhado.pk]))
    assert resp.status_code == 404


@pytest.mark.django_db
def test_comercial_nao_abre_detalhe_de_laboratorio(
    client, user_quality, user_inteligencia, user_expedicao, user_comercial
):
    """O comercial não enxerga (404) um chamado que ainda está no LABORATORIO."""
    no_lab = _em_laboratorio(user_quality, user_inteligencia, user_expedicao)
    client.force_login(user_comercial)
    resp = client.get(reverse("chamados:detalhe", args=[no_lab.pk]))
    assert resp.status_code == 404


@pytest.mark.django_db
def test_todos_os_papeis_usam_a_mesma_tela(
    client, user_quality, user_inteligencia, user_expedicao,
    user_laboratorio, user_comercial,
):
    """A tela é única (chamados:fila): todos os papéis acessam a MESMA URL com 200."""
    for usuario in (user_quality, user_inteligencia, user_expedicao,
                    user_laboratorio, user_comercial):
        client.force_login(usuario)
        resp = client.get(reverse("chamados:fila"))
        assert resp.status_code == 200


@pytest.mark.django_db
def test_expedicao_ve_so_o_que_passou_por_ela_na_fila_unica(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """A expedição vê os que estão nela e os que já passaram por ela — nunca um
    chamado que ainda não chegou à expedição."""
    em_exp = _em_expedicao(user_quality, user_inteligencia, user_expedicao=user_expedicao)
    em_lab = _em_laboratorio(user_quality, user_inteligencia, user_expedicao, user_laboratorio=user_laboratorio)
    so_encaminhado = _encaminhado_para(user_quality, user_inteligencia)

    client.force_login(user_expedicao)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    # sabotagem: Exists com setor__in=SETORES_TIMELINE (qualquer passagem) → vermelho
    assert pks == {em_exp.pk, em_lab.pk}


# --------------------------------------------------------------------------- #
# Aceite da tratativa + passagens por setor (SLA)                              #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_abertura_cria_passagem_do_quality_ja_aceita(user_quality):
    """Quem abre já é dono: a passagem do Quality nasce aceita (sem clique)."""
    from chamados.enums import Setor

    chamado = _abrir(user_quality, user_quality)
    passagem = chamado.passagens.get()
    assert passagem.setor == Setor.QUALITY
    assert passagem.aceito_em == chamado.aberto_em
    assert passagem.aceito_por == user_quality
    assert passagem.finalizado_em is None  # ainda em aberto
    assert passagem.espera.total_seconds() == 0


@pytest.mark.django_db
def test_abrir_ja_encaminhado_fecha_quality_e_abre_inteligencia(
    user_quality, user_inteligencia
):
    """RN-08: nasce em ENCAMINHADO — Quality já sai, Inteligência entra sem aceite."""
    from chamados.enums import Setor

    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    quality, intel = list(chamado.passagens.order_by("id"))
    assert quality.setor == Setor.QUALITY and quality.finalizado_em is not None
    assert quality.acao_saida == Acao.ENCAMINHAR
    assert intel.setor == Setor.INTELIGENCIA
    assert intel.aceito_em is None  # aguarda o aceite da inteligência


@pytest.mark.django_db
def test_sem_aceite_so_oferece_aceitar_e_service_recusa(
    user_quality, user_inteligencia
):
    """Antes do aceite: única ação é ACEITAR e o service recusa as demais."""
    from chamados.selectors import acoes_disponiveis

    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    assert acoes_disponiveis(user_inteligencia, chamado) == [Acao.ACEITAR_TRATATIVA]

    with pytest.raises(ValidationError):  # sem aceite não age
        services.executar(
            chamado, Acao.RESOLVER, {"procedimento_realizado": "x"}, user_inteligencia
        )
    chamado.refresh_from_db()
    assert chamado.status == Status.ENCAMINHADO  # inalterado


@pytest.mark.django_db
def test_aceitar_grava_marco_e_nao_muda_status(user_quality, user_inteligencia):
    """O aceite carimba aceito_em/por, registra o evento e NÃO muda o status."""
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    services.aceitar_tratativa(chamado, user_inteligencia)
    chamado.refresh_from_db()

    assert chamado.status == Status.ENCAMINHADO  # não mudou
    passagem = chamado.passagens.order_by("id").last()
    assert passagem.aceito_em is not None
    assert passagem.aceito_por == user_inteligencia

    evento = chamado.eventos.order_by("id").last()
    assert evento.acao == Acao.ACEITAR_TRATATIVA
    assert evento.estado_origem == evento.estado_destino == Status.ENCAMINHADO


@pytest.mark.django_db
def test_aceitar_duas_vezes_falha(user_quality, user_inteligencia):
    chamado = _encaminhado_para(user_quality, user_inteligencia)  # já aceito
    with pytest.raises(ValidationError):
        services.aceitar_tratativa(chamado, user_inteligencia)


@pytest.mark.django_db
def test_quem_nao_tem_posse_nao_aceita(
    user_quality, user_inteligencia, outro_inteligencia
):
    """Aceite segue a mesma posse das ações: intel B não aceita o chamado do A."""
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    with pytest.raises(PermissionDenied):
        services.aceitar_tratativa(chamado, outro_inteligencia)


@pytest.mark.django_db
def test_fluxo_completo_gera_uma_passagem_por_setor(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """Percorre o ciclo inteiro e confere uma passagem por setor, encadeadas."""
    from chamados.enums import Setor

    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {"finalizacao_equipamento": [
            {"numero": "EQ-001", "tratativa": "t", "custo": "SEM_CUSTO", "destino": "SUBSTITUICAO"}
        ]},
        user_comercial,
    )
    _seguir_ate_financeiro(chamado, user_expedicao)
    financeiro = _usuario_do_grupo("fin1", "financeiro")
    services.aceitar_tratativa(chamado, financeiro)
    services.executar(chamado, Acao.CONFIRMAR_ENCERRAMENTO, {}, financeiro)
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO

    passagens = list(chamado.passagens.order_by("id"))
    # A Expedição aparece duas vezes: a chegada e, no fim, o envio ao cliente.
    assert [p.setor for p in passagens] == [
        Setor.QUALITY, Setor.INTELIGENCIA, Setor.EXPEDICAO,
        Setor.LABORATORIO, Setor.COMERCIAL, Setor.RECEPCAO,
        Setor.CONFIGURACAO, Setor.EXPEDICAO, Setor.FINANCEIRO,
    ]
    # todas fechadas, com os três marcos coerentes e durações não-negativas
    for p in passagens:
        assert p.aceito_em is not None and p.finalizado_em is not None
        assert p.chegou_em <= p.aceito_em <= p.finalizado_em
        assert p.espera.total_seconds() >= 0
        assert p.trabalho.total_seconds() >= 0
        assert p.total == p.espera + p.trabalho
    # encadeamento: a saída de uma passagem é a chegada da seguinte
    for anterior, seguinte in zip(passagens, passagens[1:]):
        assert anterior.finalizado_em == seguinte.chegou_em


@pytest.mark.django_db
def test_bloquear_e_reabrir_nao_mexem_nas_passagens(user_quality):
    """Bloqueio/reabertura pausam o chamado sem trocar de setor: 1 passagem só."""
    chamado = _abrir(user_quality, user_quality)
    assert chamado.passagens.count() == 1

    services.executar(chamado, Acao.BLOQUEAR, {"motivo": "peça"}, user_quality)
    chamado.refresh_from_db()
    services.executar(chamado, Acao.REABRIR, {"motivo": "chegou"}, user_quality)
    chamado.refresh_from_db()

    assert chamado.passagens.count() == 1
    passagem = chamado.passagens.get()
    assert passagem.finalizado_em is None  # segue aberta com o mesmo dono


@pytest.mark.django_db
def test_aceitar_via_view(client, user_quality, user_inteligencia):
    """POST do aceite pela view carimba o marco e volta ao detalhe."""
    chamado = _abrir(
        user_quality, user_quality, encaminhar=True,
        procedimento_realizado="p", tratativa="t",
        responsavel_inteligencia=user_inteligencia,
    )
    client.force_login(user_inteligencia)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.ACEITAR_TRATATIVA])
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])
    passagem = chamado.passagens.order_by("id").last()
    assert passagem.aceito_por == user_inteligencia


@pytest.mark.django_db
def test_sla_nao_aparece_na_ui(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """O SLA é só para o admin: nada de passagens/tempos no detalhe do fluxo."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    client.force_login(user_laboratorio)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    for termo in ("SLA", "Passagem", "passagens", "Espera", "chegou_em"):
        assert termo not in html


# --------------------------------------------------------------------------- #
# Tratativas de contato da Expedição com o cliente                             #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_expedicao_registra_contato(user_quality, user_inteligencia, user_expedicao):
    """Registra a tentativa de contato sem mudar o status do chamado."""
    from chamados.models import ContatoExpedicao

    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    services.registrar_contato(
        chamado, user_expedicao,
        nome_contato="Maria", telefone="1133334444",
        tratativa="Sem sucesso, retorna amanhã",
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.EXPEDICAO  # não muda

    contato = ContatoExpedicao.objects.get(chamado=chamado)
    assert contato.nome_contato == "Maria"
    assert contato.telefone == "1133334444"
    assert contato.registrado_por == user_expedicao
    assert contato.codigo_rastreio == ""  # opcional


@pytest.mark.django_db
def test_varios_contatos_formam_historico(
    user_quality, user_inteligencia, user_expedicao
):
    """Cada registro é uma tentativa: o histórico acumula (mais recente primeiro)."""
    from chamados.models import ContatoExpedicao

    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    services.registrar_contato(
        chamado, user_expedicao, nome_contato="Maria", tratativa="Sem sucesso"
    )
    services.registrar_contato(
        chamado, user_expedicao, nome_contato="João",
        tratativa="Vai postar dia 20", codigo_rastreio="BR123456789BR",
    )
    contatos = ContatoExpedicao.objects.filter(chamado=chamado)
    assert contatos.count() == 2
    assert contatos.first().nome_contato == "João"  # ordering: -criado_em
    assert contatos.first().codigo_rastreio == "BR123456789BR"


@pytest.mark.django_db
def test_contato_exige_nome_e_tratativa(
    user_quality, user_inteligencia, user_expedicao
):
    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    with pytest.raises(ValidationError):
        services.registrar_contato(
            chamado, user_expedicao, nome_contato="", tratativa="x"
        )
    with pytest.raises(ValidationError):
        services.registrar_contato(
            chamado, user_expedicao, nome_contato="Maria", tratativa="   "
        )


@pytest.mark.django_db
def test_contato_exige_aceite(user_quality, user_inteligencia, user_expedicao):
    """Sem aceite, a expedição não registra contato (é trabalho do setor)."""
    chamado = _em_expedicao(user_quality, user_inteligencia)  # sem aceitar
    with pytest.raises(ValidationError):
        services.registrar_contato(
            chamado, user_expedicao, nome_contato="Maria", tratativa="Sem sucesso"
        )


@pytest.mark.django_db
def test_nao_expedicao_nao_registra_contato(
    user_quality, user_inteligencia, user_expedicao, user_comum
):
    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    with pytest.raises(PermissionDenied):
        services.registrar_contato(
            chamado, user_comum, nome_contato="Maria", tratativa="x"
        )


@pytest.mark.django_db
def test_contato_so_na_expedicao(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """Fora do estado EXPEDICAO não se registra contato (ex.: já no laboratório)."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    with pytest.raises(ValidationError):
        services.registrar_contato(
            chamado, user_laboratorio, nome_contato="Maria", tratativa="x"
        )


@pytest.mark.django_db
def test_registrar_contato_via_view(
    client, user_quality, user_inteligencia, user_expedicao
):
    """POST pela view grava o contato e volta ao detalhe."""
    from chamados.models import ContatoExpedicao

    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    client.force_login(user_expedicao)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.REGISTRAR_CONTATO]),
        {
            "nome_contato": "Andreia falou com Maria",
            "telefone": "11999998888",
            "tratativa": "Cliente vai enviar o equipamento dia 20",
            "codigo_rastreio": "BR987654321BR",
        },
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])
    contato = ContatoExpedicao.objects.get(chamado=chamado)
    assert contato.codigo_rastreio == "BR987654321BR"


@pytest.mark.django_db
def test_contatos_visiveis_da_expedicao_em_diante(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    """O histórico de contatos segue visível depois que o chamado avança."""
    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    services.registrar_contato(
        chamado, user_expedicao, nome_contato="Maria",
        tratativa="Vai postar", codigo_rastreio="BR111222333BR",
    )
    # expedição vê
    client.force_login(user_expedicao)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Tratativas de contato" in html
    assert "BR111222333BR" in html

    # avança para o laboratório: continua visível lá
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    client.force_login(user_laboratorio)
    html_lab = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Maria" in html_lab
    assert "BR111222333BR" in html_lab


@pytest.mark.django_db
def test_botao_contato_aparece_para_expedicao(
    client, user_quality, user_inteligencia, user_expedicao
):
    """A expedição vê 'Tratativas de Contato' ao lado de 'Marcar chegada'."""
    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    client.force_login(user_expedicao)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Tratativas de Contato" in html
    assert "Marcar chegada" in html
    assert "modalContatoExpedicao" in html


# --------------------------------------------------------------------------- #
# Botão "Baixar laudo" (após o Comercial aceitar a tratativa)                   #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_laudo_indisponivel_antes_do_aceite_do_comercial(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, manutencao,
):
    """Chegou no comercial COM manutenção vinculada, mas sem aceite → sem laudo."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {
            "tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}],
            "manutencao": manutencao,
        },
        user_laboratorio,
    )
    client.force_login(user_comercial)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Baixar laudo" not in html  # ainda não aceitou


@pytest.mark.django_db
def test_laudo_disponivel_apos_aceite_do_comercial(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, manutencao,
):
    """Após o aceite do comercial, o botão do laudo aparece apontando para a
    mesma URL da tela "Registro das entradas" (download_pdfmanutencao)."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {
            "tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}],
            "manutencao": manutencao,
        },
        user_laboratorio,
    )
    services.aceitar_tratativa(chamado, user_comercial)

    client.force_login(user_comercial)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Baixar laudo" in html
    assert reverse("download_pdfmanutencao", args=[manutencao.pk]) in html


@pytest.mark.django_db
def test_laudo_continua_apos_finalizar(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, manutencao,
):
    """O laudo segue disponível depois que o Comercial passa o chamado adiante
    (a passagem do comercial permanece registrada como aceita)."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {
            "tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}],
            "manutencao": manutencao,
        },
        user_laboratorio,
    )
    services.aceitar_tratativa(chamado, user_comercial)
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {"finalizacao_equipamento": [
            {"numero": "EQ-001", "tratativa": "t", "custo": "SEM_CUSTO", "destino": "DEVOLUCAO"}
        ]},
        user_comercial,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.CONFIGURACAO

    client.force_login(user_quality)  # quality vê tudo
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Baixar laudo" in html


@pytest.mark.django_db
def test_sem_manutencao_vinculada_nao_tem_laudo(
    client, user_quality, user_inteligencia, user_expedicao
):
    """Chamado sem manutenção vinculada nunca oferece o laudo."""
    chamado = _em_expedicao(
        user_quality, user_inteligencia, user_expedicao=user_expedicao
    )
    client.force_login(user_expedicao)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Baixar laudo" not in html


# --------------------------------------------------------------------------- #
# Termo de substituição (obrigatório quando há equipamento COM CUSTO)          #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_com_custo_exige_termo(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """COM CUSTO sem termo anexado → não finaliza."""
    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.FINALIZAR_COMERCIAL,
            {"finalizacao_equipamento": [
                {"numero": "EQ-001", "tratativa": "t", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
            ]},
            user_comercial,
        )
    chamado.refresh_from_db()
    assert chamado.status == Status.COMERCIAL  # inalterado


@pytest.mark.django_db
def test_sem_custo_nao_exige_termo(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """Todos SEM CUSTO → segue sem exigir o termo."""
    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {"finalizacao_equipamento": [
            {"numero": "EQ-001", "tratativa": "t", "custo": "SEM_CUSTO", "destino": "SUBSTITUICAO"}
        ]},
        user_comercial,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.RECEPCAO
    assert not chamado.termo_substituicao


@pytest.mark.django_db
def test_termo_salvo_ao_finalizar_com_custo(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {
            "finalizacao_equipamento": [
                {"numero": "EQ-001", "tratativa": "t", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
            ],
            "termo_substituicao": _pdf_falso("termo-abc.pdf"),
        },
        user_comercial,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.RECEPCAO  # substituição → recepção
    assert chamado.termo_substituicao
    assert chamado.termo_substituicao.name.endswith(".pdf")
    assert "chamados/termos/" in chamado.termo_substituicao.name


@pytest.mark.django_db
def test_form_recusa_arquivo_nao_pdf(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial,
):
    """Só PDF: um anexo de outro tipo é recusado pelo form."""
    from django.core.files.uploadedfile import SimpleUploadedFile

    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    client.force_login(user_comercial)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.FINALIZAR_COMERCIAL]),
        {
            "tratativa_0": "t", "custo_0": "COM_CUSTO", "destino_0": "SUBSTITUICAO",
            "termo_substituicao": SimpleUploadedFile(
                "termo.docx", b"nao e pdf", content_type="application/msword"
            ),
        },
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])  # erro no form
    chamado.refresh_from_db()
    assert chamado.status == Status.COMERCIAL  # não finalizou


@pytest.mark.django_db
def test_view_exige_termo_quando_ha_custo(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial,
):
    """Pela view: COM CUSTO sem anexo → volta ao detalhe com erro."""
    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    client.force_login(user_comercial)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.FINALIZAR_COMERCIAL]),
        {"tratativa_0": "t", "custo_0": "COM_CUSTO", "destino_0": "SUBSTITUICAO"},  # sem termo
    )
    assert resp.status_code == 302
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])
    chamado.refresh_from_db()
    assert chamado.status == Status.COMERCIAL


@pytest.mark.django_db
def test_termo_acessivel_a_quem_ve_o_laudo(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, manutencao,
):
    """O termo acompanha o laudo: quem tem acesso ao laudo (comercial/financeiro)
    também baixa o termo — o Financeiro precisa dos dois para cobrar o cliente."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {
            "tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}],
            "manutencao": manutencao,
        },
        user_laboratorio,
    )
    services.aceitar_tratativa(chamado, user_comercial)
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {
            "finalizacao_equipamento": [
                {"numero": "EQ-001", "tratativa": "t", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
            ],
            "termo_substituicao": _pdf_falso(),
        },
        user_comercial,
    )
    client.force_login(user_quality)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Termo de substituição" in html
    assert "Baixar laudo" in html


# --------------------------------------------------------------------------- #
# Fluxo Financeiro — com custo vai ao financeiro; sem custo encerra no comercial #
# --------------------------------------------------------------------------- #


def _requisicao_para(chamado):
    """Requisição de substituição mínima, vinculada ao chamado."""
    from requisicao.models import Requisicoes

    req = Requisicoes(nome=chamado.cliente, email="a@b.com",
                      tipo_produto=chamado.equipamentos.first().modelo,
                      numero_de_equipamentos="1", motivo="Substituição")
    req._skip_signals = True
    req.save()
    Chamado.objects.filter(pk=chamado.pk).update(requisicao=req)
    chamado.refresh_from_db()
    return req


def _seguir_ate_financeiro(chamado, user_expedicao):
    """Da saída do Comercial (RECEPCAO ou CONFIGURACAO) até o FINANCEIRO."""
    import datetime

    chamado.refresh_from_db()
    if chamado.status == Status.RECEPCAO:
        recepcao = _usuario_do_grupo("rec1", "recepcao")
        services.aceitar_tratativa(chamado, recepcao)
        _requisicao_para(chamado)
        services.executar(chamado, Acao.ENCAMINHAR_CONFIGURACAO, {}, recepcao)
    configuracao = _usuario_do_grupo("cfg1", "CONFIGURACAO")
    services.aceitar_tratativa(chamado, configuracao)
    services.executar(chamado, Acao.ENCAMINHAR_ENVIO,
                      {"tratativa": "configurado ok"}, configuracao)
    services.aceitar_tratativa(chamado, user_expedicao)
    services.executar(chamado, Acao.REGISTRAR_ENVIO, {
        "metodo_envio": "Motoboy", "data_envio": datetime.date(2026, 10, 1),
        "codigo_rastreio_envio": "",
    }, user_expedicao)
    chamado.refresh_from_db()
    return chamado


def _no_financeiro(user_quality, user_inteligencia, user_expedicao,
                   user_laboratorio, user_comercial, user_financeiro=None,
                   manutencao=None):
    """Chamado levado até FINANCEIRO: comercial finalizou COM CUSTO e SUBSTITUIÇÃO,
    e o chamado passou por Recepção, Configuração e envio pela Expedição."""
    chamado = _em_comercial(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial=user_comercial,
    )
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {
            "finalizacao_equipamento": [
                {"numero": "EQ-001", "tratativa": "orçado", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
            ],
            "termo_substituicao": _pdf_falso(),
        },
        user_comercial,
    )
    _seguir_ate_financeiro(chamado, user_expedicao)
    if user_financeiro is not None:
        services.aceitar_tratativa(chamado, user_financeiro)
    return chamado


@pytest.mark.django_db
def test_com_custo_vai_para_financeiro(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial
):
    """Havendo equipamento COM CUSTO, o comercial NÃO encerra: vai ao financeiro."""
    from chamados.enums import Setor

    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial,
    )
    assert chamado.status == Status.FINANCEIRO
    # abriu a passagem do financeiro (aguardando aceite)
    passagem = chamado.passagens.order_by("id").last()
    assert passagem.setor == Setor.FINANCEIRO
    assert passagem.aceito_em is None


@pytest.mark.django_db
@pytest.mark.parametrize("destinos, status_esperado", [
    (["SUBSTITUICAO", "DEVOLUCAO"], Status.RECEPCAO),  # algum de substituição
    (["DEVOLUCAO", "DEVOLUCAO"], Status.CONFIGURACAO),  # todos de devolução
])
def test_saida_do_comercial_depende_de_substituicao(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
    destinos, status_esperado,
):
    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio,
                            numeros=["EQ-1", "EQ-2"], user_comercial=user_comercial)
    services.executar(chamado, Acao.FINALIZAR_COMERCIAL, {"finalizacao_equipamento": [
        {"numero": n, "tratativa": "t", "custo": "SEM_CUSTO", "destino": d}
        for n, d in zip(["EQ-1", "EQ-2"], destinos)
    ]}, user_comercial)

    chamado.refresh_from_db()
    # sabotagem: rotear por custo em vez de destino → vermelho
    assert chamado.status == status_esperado


@pytest.mark.django_db
def test_financeiro_fatura_e_encerra(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_financeiro,
):
    """Financeiro aceita, informa valor + NF e o chamado é ENCERRADO."""
    from decimal import Decimal

    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial, user_financeiro=user_financeiro,
    )
    services.executar(
        chamado, Acao.FATURAR,
        {"valor_faturamento": Decimal("1250.50"), "nota_fiscal": "NF-12345"},
        user_financeiro,
    )
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO
    assert chamado.valor_faturamento == Decimal("1250.50")
    assert chamado.nota_fiscal == "NF-12345"


@pytest.mark.django_db
def test_faturar_exige_valor_e_nf(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_financeiro,
):
    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial, user_financeiro=user_financeiro,
    )
    with pytest.raises(ValidationError):  # sem NF
        services.executar(
            chamado, Acao.FATURAR,
            {"valor_faturamento": "100.00", "nota_fiscal": ""},
            user_financeiro,
        )
    chamado.refresh_from_db()
    assert chamado.status == Status.FINANCEIRO  # inalterado


@pytest.mark.django_db
def test_financeiro_precisa_aceitar_antes_de_faturar(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_financeiro,
):
    """Sem aceite, o financeiro não fatura (marco inicial do SLA do setor)."""
    from chamados.selectors import acoes_disponiveis

    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial,  # sem aceitar
    )
    assert acoes_disponiveis(user_financeiro, chamado) == [Acao.ACEITAR_TRATATIVA]
    with pytest.raises(ValidationError):
        services.executar(
            chamado, Acao.FATURAR,
            {"valor_faturamento": "100.00", "nota_fiscal": "NF-1"},
            user_financeiro,
        )


@pytest.mark.django_db
def test_nao_financeiro_nao_fatura(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_comum,
):
    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial,
    )
    with pytest.raises(PermissionDenied):
        services.executar(
            chamado, Acao.FATURAR,
            {"valor_faturamento": "100.00", "nota_fiscal": "NF-1"},
            user_comum,
        )


@pytest.mark.django_db
def test_financeiro_ve_fila_de_financeiro(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_financeiro,
):
    """O financeiro vê os chamados em FINANCEIRO; o comercial continua vendo."""
    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial,
    )
    client.force_login(user_financeiro)
    resp = client.get(reverse("chamados:fila"))
    pks = {linha["chamado"].pk for linha in resp.context["linhas"]}
    assert pks == {chamado.pk}

    client.force_login(user_comercial)
    resp = client.get(reverse("chamados:fila"))
    assert chamado.pk in {l["chamado"].pk for l in resp.context["linhas"]}


@pytest.mark.django_db
def test_faturar_via_view(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_financeiro,
):
    """POST pela view grava valor + NF e encerra o chamado."""
    from decimal import Decimal

    chamado = _no_financeiro(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio,
        user_comercial, user_financeiro=user_financeiro,
    )
    client.force_login(user_financeiro)
    resp = client.post(
        reverse("chamados:acao", args=[chamado.pk, Acao.FATURAR]),
        {"valor_faturamento": "980.00", "nota_fiscal": "NF-777"},
    )
    assert resp.status_code == 302
    chamado.refresh_from_db()
    assert chamado.status == Status.RESOLVIDO
    assert chamado.valor_faturamento == Decimal("980.00")
    assert chamado.nota_fiscal == "NF-777"


@pytest.mark.django_db
def test_financeiro_acessa_laudo_e_termo(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_financeiro, manutencao,
):
    """O financeiro precisa do laudo E do termo para cobrar do cliente."""
    chamado = _em_laboratorio(
        user_quality, user_inteligencia, user_expedicao,
        user_laboratorio=user_laboratorio,
    )
    services.executar(
        chamado, Acao.ENCAMINHAR_COMERCIAL,
        {
            "tratativas_equipamento": [{"numero": "EQ-001", "tratativa": "t"}],
            "manutencao": manutencao,
        },
        user_laboratorio,
    )
    services.aceitar_tratativa(chamado, user_comercial)
    services.executar(
        chamado, Acao.FINALIZAR_COMERCIAL,
        {
            "finalizacao_equipamento": [
                {"numero": "EQ-001", "tratativa": "t", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
            ],
            "termo_substituicao": _pdf_falso(),
        },
        user_comercial,
    )
    _seguir_ate_financeiro(chamado, user_expedicao)
    services.aceitar_tratativa(chamado, user_financeiro)

    client.force_login(user_financeiro)
    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    assert "Baixar laudo" in html
    assert "Termo de substituição" in html
    assert "Faturado" in html  # botão do modal


# --------------------------------------------------------------------------- #
# Filtro por período da linha do tempo + exportação Excel                      #
# --------------------------------------------------------------------------- #


def _entrou_em(chamado, setor, quando):
    """Reposiciona no tempo a entrada do chamado num setor (a passagem é criada
    pelo service com NOW(); os testes de período precisam de datas controladas)."""
    from chamados.models import PassagemSetor

    PassagemSetor.objects.filter(chamado=chamado, setor=setor).update(chegou_em=quando)


def _dia(dia, mes=6, ano=2026):
    import datetime

    from django.utils import timezone

    return timezone.make_aware(datetime.datetime(ano, mes, dia, 10, 0))


def _encaminhar_para_inteligencia(chamado, quality, inteligencia):
    # A passagem de abertura já nasce aceita (aceito_em = aberto_em): quem abriu
    # não precisa clicar "aceitar", então aqui só encaminha.
    services.executar(
        chamado,
        Acao.ENCAMINHAR,
        {
            "procedimento_realizado": "p",
            "tratativa": "t",
            "responsavel_inteligencia": inteligencia,
        },
        quality,
    )


@pytest.mark.django_db
def test_filtro_por_setor_pega_quem_ja_saiu_do_setor(
    user_quality, user_inteligencia, user_expedicao
):
    """O recorte é a ENTRADA no setor, não o status atual: um chamado que passou
    pela Expedição e já seguiu para o Laboratório continua no resultado."""
    from chamados.selectors import listar_fila

    passou = _abrir(user_quality, user_quality)
    _encaminhar_para_inteligencia(passou, user_quality, user_inteligencia)
    services.aceitar_tratativa(passou, user_inteligencia)
    services.executar(
        passou,
        Acao.ENCAMINHAR_EXPEDICAO,
        {"procedimento_realizado": "p", "tratativa": "t"},
        user_inteligencia,
    )
    _entrou_em(passou, Setor.EXPEDICAO, _dia(10))
    # Sai da Expedição: o status muda, mas a passagem (e a entrada) permanece.
    services.aceitar_tratativa(passou, user_expedicao)
    services.executar(passou, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    passou.refresh_from_db()
    assert passou.status == Status.LABORATORIO  # já NÃO está na expedição

    nunca = _abrir(user_quality, user_quality)  # nunca passou pela Expedição

    resultado = listar_fila(
        user_quality,
        setor=Setor.EXPEDICAO,
        data_de=_dia(1).date(),
        data_ate=_dia(30).date(),
    )
    protocolos = {c.protocolo for c in resultado}
    assert passou.protocolo in protocolos
    assert nunca.protocolo not in protocolos


@pytest.mark.django_db
@pytest.mark.parametrize(
    "de, ate, esperado",
    [
        (10, 20, True),   # dentro da janela
        (15, 15, True),   # janela de um dia só, no dia da entrada — inclusiva
        (16, 30, False),  # janela começa depois da entrada
        (1, 14, False),   # janela termina antes da entrada
    ],
)
def test_periodo_recorta_pela_entrada_no_setor(user_quality, de, ate, esperado):
    """As duas pontas do período são INCLUSIVAS e comparadas por data."""
    from chamados.selectors import listar_fila

    chamado = _abrir(user_quality, user_quality)
    _entrou_em(chamado, Setor.QUALITY, _dia(15))

    resultado = listar_fila(
        user_quality,
        setor=Setor.QUALITY,
        data_de=_dia(de).date(),
        data_ate=_dia(ate).date(),
    )
    assert (chamado.protocolo in {c.protocolo for c in resultado}) is esperado


@pytest.mark.django_db
def test_sem_setor_o_periodo_recorta_a_abertura(user_quality):
    """Sem setor escolhido, "de tal data até tal data" é sobre a ABERTURA."""
    from chamados.models import Chamado
    from chamados.selectors import listar_fila

    antigo = _abrir(user_quality, user_quality)
    recente = _abrir(user_quality, user_quality)
    Chamado.objects.filter(pk=antigo.pk).update(aberto_em=_dia(5))
    Chamado.objects.filter(pk=recente.pk).update(aberto_em=_dia(25))

    resultado = listar_fila(
        user_quality, setor=None, data_de=_dia(20).date(), data_ate=_dia(30).date()
    )
    protocolos = {c.protocolo for c in resultado}
    assert recente.protocolo in protocolos
    assert antigo.protocolo not in protocolos


@pytest.mark.django_db
def test_fila_nao_faz_query_por_linha(user_quality, django_assert_num_queries):
    """A linha do tempo é anotada por Subquery: a contagem de queries é CONSTANTE,
    não cresce com o número de chamados (mandamentos 1, 2 e 10)."""
    from chamados.selectors import SETORES_TIMELINE, campo_entrada, listar_fila

    def _consumir():
        return [
            [getattr(c, campo_entrada(s), None) for s in SETORES_TIMELINE]
            for c in listar_fila(user_quality)
        ]

    # 3 queries: duas checagens de grupo (visibilidade e corte de visibilidade)
    # + a listagem com as subqueries de entrada embutidas. Nenhuma depende do nº
    # de linhas.
    for _ in range(3):
        _abrir(user_quality, user_quality)
    with django_assert_num_queries(3):
        assert len(_consumir()) == 3

    for _ in range(4):
        _abrir(user_quality, user_quality)
    with django_assert_num_queries(3):  # mesma contagem, mais que o dobro de linhas
        assert len(_consumir()) == 7


@pytest.mark.django_db
def test_exportar_devolve_xlsx_com_as_linhas_filtradas(client, user_quality):
    """O arquivo sai com o cabeçalho da linha do tempo, só as linhas do filtro, e
    a data de entrada no setor como DATETIME (não texto — texto não ordena)."""
    import io

    from openpyxl import load_workbook

    dentro = _abrir(user_quality, user_quality)
    fora = _abrir(user_quality, user_quality)
    _entrou_em(dentro, Setor.QUALITY, _dia(15))
    _entrou_em(fora, Setor.QUALITY, _dia(25))

    client.force_login(user_quality)
    resposta = client.get(
        reverse("chamados:exportar"),
        {"setor": Setor.QUALITY, "data_de": "2026-06-10", "data_ate": "2026-06-20"},
    )

    assert resposta.status_code == 200
    assert resposta["Content-Type"].endswith("spreadsheetml.sheet")
    assert "chamados_setor-quality_de-10-06-2026_ate-20-06-2026.xlsx" in (
        resposta["Content-Disposition"]
    )

    ws = load_workbook(io.BytesIO(resposta.content)).active
    cabecalho = [c.value for c in ws[1]]
    assert cabecalho[0] == "Protocolo"
    assert "Entrou em Expedição" in cabecalho

    linhas = list(ws.iter_rows(min_row=2, values_only=True))
    assert [linha[0] for linha in linhas] == [dentro.protocolo]  # `fora` ficou fora

    assert linhas[0][cabecalho.index("Modelo")] == "Rastreador GT06"
    valor = linhas[0][cabecalho.index("Entrou em Quality")]
    assert hasattr(valor, "year"), "data deve ser datetime, nao string"
    assert (valor.year, valor.month, valor.day) == (2026, 6, 15)


@pytest.mark.django_db
def test_exportacao_respeita_a_fronteira_de_visibilidade(
    client, user_quality, user_inteligencia, outro_inteligencia
):
    """Inteligência não exporta chamado que não enxerga na tela — o gate é o
    mesmo queryset da fila, não um filtro só de UI."""
    import io

    from openpyxl import load_workbook

    meu = _abrir(user_quality, user_quality)
    _encaminhar_para_inteligencia(meu, user_quality, user_inteligencia)
    alheio = _abrir(user_quality, user_quality)
    _encaminhar_para_inteligencia(alheio, user_quality, outro_inteligencia)

    client.force_login(user_inteligencia)
    resposta = client.get(reverse("chamados:exportar"))

    ws = load_workbook(io.BytesIO(resposta.content)).active
    protocolos = {linha[0] for linha in ws.iter_rows(min_row=2, values_only=True)}
    assert meu.protocolo in protocolos
    assert alheio.protocolo not in protocolos


# --------------------------------------------------------------------------- #
# Horizonte: quem já passou o chamado adiante vê só até o seu setor            #
# --------------------------------------------------------------------------- #

_DO_LAB = ["Tratativas por equipamento", "Encaminhar para comercial"]
_DO_COMERCIAL = ["orçado", "Com custo", "Finalizar chamado", "Baixar laudo"]


@pytest.mark.django_db
@pytest.mark.parametrize(
    "papel, visiveis, ocultos",
    [
        ("user_expedicao", ["Marcar chegada", "Entrada de equipamento #"], _DO_LAB + _DO_COMERCIAL),
        ("user_laboratorio", _DO_LAB, _DO_COMERCIAL),
        ("user_comercial", _DO_LAB + _DO_COMERCIAL, []),
        ("user_inteligencia", _DO_LAB + _DO_COMERCIAL, []),  # Inteligência vê tudo
    ],
)
def test_detalhe_mostra_so_ate_o_setor_do_usuario(
    request, client, user_quality, user_inteligencia, user_expedicao,
    user_laboratorio, user_comercial, manutencao, papel, visiveis, ocultos,
):
    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio,
                            user_comercial=user_comercial)
    services.executar(chamado, Acao.FINALIZAR_COMERCIAL, {
        "finalizacao_equipamento": [
            {"numero": "EQ-001", "tratativa": "orçado", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"}
        ],
        "termo_substituicao": _pdf_falso(),
    }, user_comercial)  # → RECEPCAO: a Expedição ainda não voltou para o envio
    Chamado.objects.filter(pk=chamado.pk).update(manutencao=manutencao)
    client.force_login(request.getfixturevalue(papel))

    html = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()

    # sabotagem: corte_de_visibilidade devolver None para todos → vermelho
    # sabotagem: tirar "and ve.comercial" da tratativa comercial no detalhe → vermelho
    assert [t for t in visiveis if t not in html] == []
    assert [t for t in ocultos if t in html] == []


@pytest.mark.django_db
def test_ver_chamado_passado_nao_da_acao_sobre_ele(
    client, user_quality, user_inteligencia, user_expedicao
):
    """Visibilidade ampliada não é posse: a expedição não age no chamado que já
    foi para o laboratório."""
    from chamados.models import ContatoExpedicao

    chamado = _em_laboratorio(user_quality, user_inteligencia, user_expedicao)
    client.force_login(user_expedicao)

    client.post(reverse("chamados:acao", args=[chamado.pk, Acao.REGISTRAR_CONTATO]),
                {"nome_contato": "Fulano", "tratativa": "ligou"})
    client.post(reverse("chamados:acao", args=[chamado.pk, Acao.MARCAR_CHEGADA]))

    chamado.refresh_from_db()
    assert chamado.status == Status.LABORATORIO
    assert not ContatoExpedicao.objects.filter(chamado=chamado).exists()


@pytest.mark.django_db
def test_linha_do_tempo_da_fila_e_do_excel_para_no_setor_do_usuario(
    client, user_quality, user_inteligencia, user_expedicao
):
    import io

    from openpyxl import load_workbook

    _em_comercial(user_quality, user_inteligencia, user_expedicao, _usuario_do_grupo("lab9", "laboratorio"))
    client.force_login(user_expedicao)

    resp = client.get(reverse("chamados:fila"))
    setores = [s["label"] for s in resp.context["setores_timeline"]]
    entradas = dict(zip(setores, resp.context["linhas"][0]["entradas"]))
    ws = load_workbook(io.BytesIO(client.get(reverse("chamados:exportar")).content)).active
    linha = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))

    # A Expedição saiu ao marcar a chegada (a entrada no Laboratório é esse mesmo
    # instante): não vê quando o chamado chegou ao Comercial.
    # sabotagem: entradas_visiveis devolver as entradas sem corte → vermelho
    assert entradas["Expedição"] is not None and entradas["Comercial"] is None
    assert linha["Entrou em Expedição"] is not None and linha["Entrou em Comercial"] is None


# --------------------------------------------------------------------------- #
# Laboratório: aceitar → registrar a manutenção                               #
# --------------------------------------------------------------------------- #


def _lab_pode_editar_entrada(user):
    from django.contrib.auth.models import Permission

    user.user_permissions.add(*Permission.objects.filter(
        codename__in=["add_registrodemanutencao", "change_registrodemanutencao"]
    ))


def _entrada_do_chamado(chamado):
    """Entrada vinculada ao chamado, com os nºs dele (como a Expedição grava)."""
    from registrodemanutencao.models import ItemEntrada, registrodemanutencao

    entrada = registrodemanutencao.objects.create(nome=chamado.cliente, status="Pendente")
    for e in chamado.equipamentos.all():
        ItemEntrada.objects.create(registro=entrada, tipo_produto=e.modelo,
                                   numero_equipamento=e.numero, quantidade=1)
    Chamado.objects.filter(pk=chamado.pk).update(manutencao=entrada)
    return entrada


@pytest.mark.django_db
@pytest.mark.parametrize("com_entrada", [True, False])
def test_lab_aceita_e_vai_registrar_a_manutencao(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio, com_entrada
):
    chamado = _em_laboratorio(user_quality, user_inteligencia, user_expedicao)
    entrada = _entrada_do_chamado(chamado) if com_entrada else None
    client.force_login(user_laboratorio)

    resp = client.post(reverse("chamados:acao", args=[chamado.pk, Acao.ACEITAR_TRATATIVA]))

    esperado = (reverse("FormulariosUpdateView", args=[entrada.pk]) if com_entrada
                else reverse("FormulariosCreateView"))
    assert resp.url == f"{esperado}?chamado={chamado.pk}"


@pytest.mark.django_db
def test_laudo_vem_com_uma_linha_por_equipamento_e_ignora_linha_nao_preenchida(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    chamado, isca_4g, isca_2g = _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao)
    entrada = _entrada_do_chamado(chamado)
    _lab_pode_editar_entrada(user_laboratorio)
    client.force_login(user_laboratorio)
    url = f"{reverse('FormulariosUpdateView', args=[entrada.pk])}?chamado={chamado.pk}"

    formset = client.get(url).context["imagens_formset"]
    assert [f.initial.get("id_equipamento") for f in formset.extra_forms] == ["EQ-1", "EQ-2", "EQ-3"]

    dados = {"nome": entrada.nome_id, "tipo_entrada": "Manutenção", "status": "Pendente",
             "observacoes": "", "chamado": chamado.pk, "imagens-TOTAL_FORMS": "3",
             "imagens-INITIAL_FORMS": "0", "imagens-MIN_NUM_FORMS": "0", "imagens-MAX_NUM_FORMS": "1000"}
    for i, numero in enumerate(["EQ-1", "EQ-2", "EQ-3"]):
        dados.update({f"imagens-{i}-id_equipamento": numero, f"imagens-{i}-tipo_problema": "",
                      f"imagens-{i}-faturamento": "", f"imagens-{i}-observacao2": ""})
    dados["imagens-1-tipo_problema"] = "Oxidação"  # só EQ-2 foi preenchido

    resp = client.post(url, dados)

    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])
    # sabotagem: montar o formset do POST sem o initial do chamado → vermelho
    assert list(entrada.imagens.values_list("id_equipamento", flat=True)) == ["EQ-2"]


@pytest.mark.django_db
def test_lab_cria_entrada_e_segue_para_o_laudo(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio
):
    chamado, isca_4g, isca_2g = _chegou_no_laboratorio(user_quality, user_inteligencia, user_expedicao)
    _lab_pode_editar_entrada(user_laboratorio)
    client.force_login(user_laboratorio)

    resp = client.post(reverse("FormulariosCreateView"),
                       _post_entrada(chamado, (isca_4g, "EQ-1 EQ-3"), (isca_2g, "EQ-2")))

    chamado.refresh_from_db()
    assert resp.url == f"{reverse('FormulariosUpdateView', args=[chamado.manutencao_id])}?chamado={chamado.pk}"


# --------------------------------------------------------------------------- #
# Comercial: substituição ou devolução por equipamento                        #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
@pytest.mark.parametrize("destino, gravou", [("DEVOLUCAO", "DEVOLUCAO"), ("", None)])
def test_comercial_informa_substituicao_ou_devolucao(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
    destino, gravou,
):
    from chamados.models import TratativaEquipamento

    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio,
                            user_comercial=user_comercial)
    dados = {"finalizacao_equipamento": [
        {"numero": "EQ-001", "tratativa": "t", "custo": "SEM_CUSTO", "destino": destino}
    ]}

    if gravou is None:
        with pytest.raises(ValidationError):
            services.executar(chamado, Acao.FINALIZAR_COMERCIAL, dados, user_comercial)
    else:
        services.executar(chamado, Acao.FINALIZAR_COMERCIAL, dados, user_comercial)
    linha = TratativaEquipamento.objects.get(chamado=chamado)
    assert (linha.destino or None) == gravou


# --------------------------------------------------------------------------- #
# Abertura: customização e contrato por bloco de modelo                       #
# --------------------------------------------------------------------------- #


@pytest.mark.django_db
def test_abertura_grava_customizacao_e_contrato_e_a_entrada_herda(
    client, user_quality, user_inteligencia, user_expedicao, cliente
):
    isca_4g, isca_2g = Produto.objects.create(nome="Isca 4G"), Produto.objects.create(nome="Isca 2G")
    client.force_login(user_quality)
    client.post(reverse("chamados:abrir"), {
        "cliente": cliente.pk, "categoria": "HARDWARE", "problema_relatado": "Falha",
        "contato_nome": "Contato", "contato_meio": "TELEFONE",
        **_blocos((isca_4g.pk, ["EQ-1"], "Termo branco", "Retornavel"),
                  (isca_2g.pk, ["EQ-2"], "Caixa de papelão", "Descartavel")),
    })
    chamado = Chamado.objects.get(cliente=cliente)

    blocos = services.dados_iniciais_entrada(chamado)["blocos"]

    assert [(b["tipo_produto"], b["customizacao"], b["tipo_contrato"]) for b in blocos] == [
        (str(isca_4g.pk), "Termo branco", "Retornavel"),
        (str(isca_2g.pk), "Caixa de papelão", "Descartavel"),
    ]


@pytest.mark.django_db
@pytest.mark.parametrize("customizacao, contrato, erro", [
    ("", "Retornavel", "Selecione a customização do modelo Rastreador GT06."),
    ("Termo branco", "", "Selecione o tipo de contrato do modelo Rastreador GT06."),
])
def test_abertura_exige_customizacao_e_contrato(
    client, user_quality, cliente, produto, customizacao, contrato, erro
):
    client.force_login(user_quality)
    resp = client.post(reverse("chamados:abrir"), {
        "cliente": cliente.pk, "categoria": "HARDWARE", "problema_relatado": "Falha",
        "contato_nome": "Contato", "contato_meio": "TELEFONE",
        **_blocos((produto.pk, ["EQ-1"], customizacao, contrato)),
    })

    assert erro in resp.context["form"].errors["equipamentos"]
    assert not Chamado.objects.exists()


# --------------------------------------------------------------------------- #
# Depois do Comercial: Recepção, Configuração, envio e Financeiro             #
# --------------------------------------------------------------------------- #


def _na_recepcao(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
                 custo="SEM_CUSTO"):
    chamado = _em_comercial(user_quality, user_inteligencia, user_expedicao, user_laboratorio,
                            user_comercial=user_comercial)
    dados = {"finalizacao_equipamento": [
        {"numero": "EQ-001", "tratativa": "orçado", "custo": custo, "destino": "SUBSTITUICAO"}
    ]}
    if custo == "COM_CUSTO":
        dados["termo_substituicao"] = _pdf_falso()
    services.executar(chamado, Acao.FINALIZAR_COMERCIAL, dados, user_comercial)
    chamado.refresh_from_db()
    return chamado


@pytest.mark.django_db
def test_recepcao_so_encaminha_com_requisicao_vinculada(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial, user_recepcao
):
    from chamados.selectors import acoes_disponiveis

    chamado = _na_recepcao(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    services.aceitar_tratativa(chamado, user_recepcao)

    assert Acao.ENCAMINHAR_CONFIGURACAO not in acoes_disponiveis(user_recepcao, chamado)
    # sabotagem: remover a checagem de requisicao_id em executar → vermelho
    with pytest.raises(ValidationError):
        services.executar(chamado, Acao.ENCAMINHAR_CONFIGURACAO, {}, user_recepcao)

    _requisicao_para(chamado)
    services.executar(chamado, Acao.ENCAMINHAR_CONFIGURACAO, {}, user_recepcao)
    chamado.refresh_from_db()
    assert chamado.status == Status.CONFIGURACAO


@pytest.mark.django_db
@pytest.mark.parametrize("etapa, acao, dados, papel_errado", [
    ("recepcao", Acao.ENCAMINHAR_CONFIGURACAO, {}, "user_configuracao"),
    ("configuracao", Acao.ENCAMINHAR_ENVIO, {"tratativa": "x"}, "user_recepcao"),
    ("envio", Acao.REGISTRAR_ENVIO, {"metodo_envio": "Motoboy", "data_envio": "2026-10-01"}, "user_configuracao"),
])
def test_cada_etapa_nova_so_aceita_o_proprio_setor(
    request, user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
    user_recepcao, user_configuracao, etapa, acao, dados, papel_errado,
):
    chamado = _na_recepcao(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    _requisicao_para(chamado)
    if etapa in ("configuracao", "envio"):
        services.aceitar_tratativa(chamado, user_recepcao)
        services.executar(chamado, Acao.ENCAMINHAR_CONFIGURACAO, {}, user_recepcao)
    if etapa == "envio":
        services.aceitar_tratativa(chamado, user_configuracao)
        services.executar(chamado, Acao.ENCAMINHAR_ENVIO, {"tratativa": "ok"}, user_configuracao)

    # sabotagem: pode_agir devolver True para RECEPCAO/CONFIGURACAO/ENVIO → vermelho
    with pytest.raises(PermissionDenied):
        services.aceitar_tratativa(chamado, request.getfixturevalue(papel_errado))
    with pytest.raises(PermissionDenied):
        services.executar(chamado, acao, dados, request.getfixturevalue(papel_errado))


@pytest.mark.django_db
def test_envio_registra_metodo_data_e_rastreio(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
    user_recepcao, user_configuracao,
):
    chamado = _na_recepcao(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    _requisicao_para(chamado)
    services.aceitar_tratativa(chamado, user_recepcao)
    services.executar(chamado, Acao.ENCAMINHAR_CONFIGURACAO, {}, user_recepcao)
    services.aceitar_tratativa(chamado, user_configuracao)
    services.executar(chamado, Acao.ENCAMINHAR_ENVIO, {"tratativa": "ok"}, user_configuracao)
    services.aceitar_tratativa(chamado, user_expedicao)
    client.force_login(user_expedicao)

    sem_data = client.post(reverse("chamados:acao", args=[chamado.pk, Acao.REGISTRAR_ENVIO]),
                           {"metodo_envio": "Correio", "codigo_rastreio_envio": "BR123"})
    chamado.refresh_from_db()
    assert chamado.status == Status.ENVIO  # data obrigatória
    client.post(reverse("chamados:acao", args=[chamado.pk, Acao.REGISTRAR_ENVIO]),
                {"metodo_envio": "Correio", "data_envio": "2026-10-01", "codigo_rastreio_envio": "BR123"})

    chamado.refresh_from_db()
    assert sem_data.status_code == 302
    assert (chamado.status, chamado.metodo_envio, str(chamado.data_envio), chamado.codigo_rastreio_envio) == (
        Status.FINANCEIRO, "Correio", "2026-10-01", "BR123")


@pytest.mark.django_db
@pytest.mark.parametrize("custo, oferecida, recusada", [
    ("SEM_CUSTO", Acao.CONFIRMAR_ENCERRAMENTO, Acao.FATURAR),
    ("COM_CUSTO", Acao.FATURAR, Acao.CONFIRMAR_ENCERRAMENTO),
])
def test_financeiro_fatura_com_custo_e_so_confirma_sem_custo(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
    user_financeiro, custo, oferecida, recusada,
):
    from chamados.selectors import acoes_disponiveis

    chamado = _na_recepcao(user_quality, user_inteligencia, user_expedicao, user_laboratorio,
                           user_comercial, custo=custo)
    _seguir_ate_financeiro(chamado, user_expedicao)
    services.aceitar_tratativa(chamado, user_financeiro)

    acoes = acoes_disponiveis(user_financeiro, chamado)
    assert oferecida in acoes and recusada not in acoes
    # sabotagem: remover a checagem de custo em executar → vermelho
    with pytest.raises(ValidationError):
        services.executar(chamado, recusada,
                          {"valor_faturamento": "10", "nota_fiscal": "NF1"}, user_financeiro)


@pytest.mark.django_db
def test_expedicao_no_envio_ve_o_que_veio_depois_da_chegada(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
    user_recepcao, user_configuracao,
):
    """De volta para o envio, a Expedição vê laboratório e comercial (precisa saber
    o que mandar); a Recepção, que saiu antes, não vê a Configuração nem o envio."""
    chamado = _na_recepcao(user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    _seguir_ate_financeiro(chamado, user_expedicao)

    client.force_login(user_expedicao)
    html_exp = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()
    client.force_login(user_recepcao)
    html_rec = client.get(reverse("chamados:detalhe", args=[chamado.pk])).content.decode()

    assert "orçado" in html_exp and "configurado ok" in html_exp and "Envio ao cliente" in html_exp
    assert "orçado" in html_rec
    assert "configurado ok" not in html_rec and "Envio ao cliente" not in html_rec


# --------------------------------------------------------------------------- #
# Recepção: aceitar → requisição de substituição preenchida                   #
# --------------------------------------------------------------------------- #


def _recepcao_pode_criar_requisicao(user):
    from django.contrib.auth.models import Permission

    user.user_permissions.add(*Permission.objects.filter(
        codename__in=["add_requisicoes", "view_requisicoes"]
    ))


def _na_recepcao_com_modelos(user_quality, user_inteligencia, user_expedicao,
                             user_laboratorio, user_comercial):
    """EQ-1 (4G, Termo branco + Imã, Retornavel) e EQ-2 (2G) de SUBSTITUIÇÃO —
    EQ-1 com custo —, EQ-3 (4G) de DEVOLUÇÃO. Chamado na RECEPÇÃO."""
    isca_4g, isca_2g = Produto.objects.create(nome="Isca 4G"), Produto.objects.create(nome="Isca 2G")
    chamado = _abrir(user_quality, user_quality, equipamentos_override=[
        ("EQ-1", isca_4g, "Termo branco + Imã", "Retornavel"),
        ("EQ-2", isca_2g, "Sem customização", "Retornavel"),
        ("EQ-3", isca_4g, "Termo branco + Imã", "Retornavel"),
    ], encaminhar=True, procedimento_realizado="p", tratativa="t",
       responsavel_inteligencia=user_inteligencia)
    services.aceitar_tratativa(chamado, user_inteligencia)
    services.executar(chamado, Acao.ENCAMINHAR_EXPEDICAO,
                      {"procedimento_realizado": "p", "tratativa": "t"}, user_inteligencia)
    services.aceitar_tratativa(chamado, user_expedicao)
    services.executar(chamado, Acao.MARCAR_CHEGADA, {}, user_expedicao)
    services.aceitar_tratativa(chamado, user_laboratorio)
    services.executar(chamado, Acao.ENCAMINHAR_COMERCIAL, {"tratativas_equipamento": [
        {"numero": n, "tratativa": "lab"} for n in ("EQ-1", "EQ-2", "EQ-3")
    ]}, user_laboratorio)
    services.aceitar_tratativa(chamado, user_comercial)
    services.executar(chamado, Acao.FINALIZAR_COMERCIAL, {
        "finalizacao_equipamento": [
            {"numero": "EQ-1", "tratativa": "c", "custo": "COM_CUSTO", "destino": "SUBSTITUICAO"},
            {"numero": "EQ-2", "tratativa": "c", "custo": "SEM_CUSTO", "destino": "SUBSTITUICAO"},
            {"numero": "EQ-3", "tratativa": "c", "custo": "SEM_CUSTO", "destino": "DEVOLUCAO"},
        ],
        "termo_substituicao": _pdf_falso(),
    }, user_comercial)
    chamado.refresh_from_db()
    return chamado, isca_4g, isca_2g


@pytest.mark.django_db
def test_recepcao_aceita_e_vai_abrir_a_requisicao(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_recepcao,
):
    chamado, *_ = _na_recepcao_com_modelos(user_quality, user_inteligencia, user_expedicao,
                                           user_laboratorio, user_comercial)
    client.force_login(user_recepcao)

    resp = client.post(reverse("chamados:acao", args=[chamado.pk, Acao.ACEITAR_TRATATIVA]))

    assert resp.url == f"{reverse('requisicoescrateview')}?chamado={chamado.pk}"


@pytest.mark.django_db
def test_requisicao_vem_so_com_os_equipamentos_de_substituicao(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_recepcao,
):
    chamado, isca_4g, isca_2g = _na_recepcao_com_modelos(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    _recepcao_pode_criar_requisicao(user_recepcao)
    client.force_login(user_recepcao)

    form = client.get(f"{reverse('requisicoescrateview')}?chamado={chamado.pk}").context["form"]

    assert (form.initial["nome"], form.initial["contrato"], form.initial["motivo"], form.initial["tipo_fatura"]) == (
        chamado.cliente_id, "Retornavel", "Substituição", "Com Custo")
    # EQ-3 (devolução) fica de fora; customização casada com o vocabulário da requisição.
    assert [(b["tipo_produto"], b["quantidade"], b["customizacao"], b["numeros"])
            for b in form.blocos_itens()] == [
        (str(isca_4g.pk), "1", "Termo branco + imã", "EQ-1"),
        (str(isca_2g.pk), "1", "Sem custumização", "EQ-2"),
    ]


@pytest.mark.django_db
def test_salvar_requisicao_vincula_ao_chamado_e_libera_a_configuracao(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_recepcao,
):
    from chamados.selectors import acoes_disponiveis

    chamado, isca_4g, isca_2g = _na_recepcao_com_modelos(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    services.aceitar_tratativa(chamado, user_recepcao)
    _recepcao_pode_criar_requisicao(user_recepcao)
    client.force_login(user_recepcao)

    resp = client.post(reverse("requisicoescrateview"), {
        "chamado": chamado.pk, "nome": chamado.cliente_id, "email": "c@acme.com",
        "contrato": "Retornavel", "motivo": "Substituição", "status": "Pendente", "taxa_envio": "0",
        "item": ["0", "1"],
        "item_0_tipo_produto": isca_4g.pk, "item_0_quantidade": "1", "item_0_valor_unitario": "0",
        "item_1_tipo_produto": isca_2g.pk, "item_1_quantidade": "1", "item_1_valor_unitario": "0",
    })

    chamado.refresh_from_db()
    assert resp.url == reverse("chamados:detalhe", args=[chamado.pk])
    assert chamado.requisicao.itens.count() == 2
    assert Acao.ENCAMINHAR_CONFIGURACAO in acoes_disponiveis(user_recepcao, chamado)


@pytest.mark.django_db
def test_vinculo_falho_desfaz_a_requisicao(
    user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
):
    """Chamado já vinculado por outra aba: a requisição nova não fica órfã."""
    from requisicao.models import ItemRequisicao, Requisicoes
    from requisicao.services import criar_requisicao

    chamado, isca_4g, _ = _na_recepcao_com_modelos(
        user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial)
    _requisicao_para(chamado)
    antes = Requisicoes.objects.count()

    with pytest.raises(ValidationError):
        criar_requisicao(Requisicoes(nome=chamado.cliente, email="a@b.com"),
                         [ItemRequisicao(tipo_produto=isca_4g, quantidade=1)], chamado=chamado)
    assert Requisicoes.objects.count() == antes


@pytest.mark.django_db
def test_requisicao_nao_vem_do_chamado_para_quem_nao_e_recepcao(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio, user_comercial,
):
    chamado, *_ = _na_recepcao_com_modelos(user_quality, user_inteligencia, user_expedicao,
                                           user_laboratorio, user_comercial)
    _recepcao_pode_criar_requisicao(user_comercial)
    client.force_login(user_comercial)

    resp = client.get(f"{reverse('requisicoescrateview')}?chamado={chamado.pk}")

    assert resp.context["chamado"] is None
    assert "nome" not in resp.context["form"].initial



@pytest.mark.django_db
@pytest.mark.parametrize("inicio", [__import__("datetime").date(2025, 1, 30), None])
def test_requisicao_preenche_com_inicio_de_contrato_do_cliente(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_recepcao, inicio,
):
    """O início de contrato do cadastro é DateField (antes era tratado como texto)."""
    chamado, *_ = _na_recepcao_com_modelos(user_quality, user_inteligencia, user_expedicao,
                                           user_laboratorio, user_comercial)
    Clientes.objects.filter(pk=chamado.cliente_id).update(inicio_de_contrato=inicio)
    _recepcao_pode_criar_requisicao(user_recepcao)
    client.force_login(user_recepcao)

    resp = client.get(f"{reverse('requisicoescrateview')}?chamado={chamado.pk}")

    assert resp.context["form"].initial["inicio_de_contrato"] == inicio



@pytest.mark.django_db
def test_recepcao_aceita_sem_requisicao_ve_o_botao_de_criar(
    client, user_quality, user_inteligencia, user_expedicao, user_laboratorio,
    user_comercial, user_recepcao,
):
    """Aceite feito e requisição não salva: o detalhe oferece abrir a requisição
    (antes o chamado ficava sem nenhuma ação para a Recepção)."""
    chamado, *_ = _na_recepcao_com_modelos(user_quality, user_inteligencia, user_expedicao,
                                           user_laboratorio, user_comercial)
    services.aceitar_tratativa(chamado, user_recepcao)
    client.force_login(user_recepcao)
    url = reverse("chamados:detalhe", args=[chamado.pk])
    link = f"{reverse('requisicoescrateview')}?chamado={chamado.pk}"

    assert link in client.get(url).content.decode()
    _requisicao_para(chamado)
    assert link not in client.get(url).content.decode()  # com requisição, some
