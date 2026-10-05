"""Planilhas para fora do sistema.

`saldo_agentes_xlsx` responde "quais iscas estão com quais agentes, por
região": uma linha por agente × modelo, ordenada por UF, com os
identificadores das unidades — é o que se confere em campo.
"""
from io import BytesIO

from django.db.models import Exists
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from iscas.enums import TipoCustodia
from iscas.models.custodia import Unidade
from iscas.services.saldo import _reserva_ativa_subquery

SEM_UF = "Sem UF"

#: Limite de caracteres de uma célula no Excel. Agente com milhares de
#: unidades estouraria a coluna de IDs e o arquivo abriria com erro.
_LIMITE_CELULA = 32_000

_CABECALHO = PatternFill(start_color="1F2933", end_color="1F2933", fill_type="solid")


def linhas_saldo_agentes(busca: str = "") -> list[dict]:
    """Uma linha por (agente, modelo) com quantidade e identificadores.

    Uma consulta só, agrupada em memória: o agrupamento por agente e modelo
    precisa dos identificadores, que `values().annotate()` não concatena de
    forma portável entre SQLite e Postgres.
    """
    unidades = (
        Unidade.objects.filter(custodia_atual__tipo=TipoCustodia.AGENTE)
        .annotate(reservada=Exists(_reserva_ativa_subquery()))
        .select_related("modelo", "custodia_atual__agente")
        .order_by("custodia_atual__agente_id", "modelo__nome", "identificador")
    )
    if busca:
        unidades = unidades.filter(custodia_atual__agente__nome__icontains=busca)

    grupos = {}
    for unidade in unidades.iterator(chunk_size=1000):
        agente = unidade.custodia_atual.agente
        chave = (agente.pk, unidade.modelo_id)
        linha = grupos.get(chave)
        if linha is None:
            linha = grupos[chave] = {
                "uf": (agente.uf or "").upper() or SEM_UF,
                "cidade": agente.cidade,
                "agente": agente.nome + ("" if agente.is_active else " (inativo)"),
                "telefone": agente.telefone,
                "modelo": unidade.modelo.nome,
                "tipo": unidade.modelo.get_tipo_display(),
                "quantidade": 0,
                "disponivel": 0,
                "reservado": 0,
                "ids": [],
            }
        linha["quantidade"] += 1
        linha["reservado" if unidade.reservada else "disponivel"] += 1
        linha["ids"].append(unidade.identificador)

    # "Sem UF" no fim: é pendência de cadastro, não uma região.
    return sorted(
        grupos.values(),
        key=lambda l: (l["uf"] == SEM_UF, l["uf"], l["agente"].lower(), l["modelo"]),
    )


def saldo_agentes_xlsx(busca: str = "") -> bytes:
    """Planilha com a aba detalhada e um resumo por UF."""
    linhas = linhas_saldo_agentes(busca)

    livro = Workbook()
    detalhe = livro.active
    detalhe.title = "Iscas com agentes"
    _escrever(
        detalhe,
        ["UF", "Cidade", "Agente", "Telefone", "Modelo", "Tipo",
         "Quantidade", "Disponível", "Reservado", "IDs"],
        (
            [l["uf"], l["cidade"], l["agente"], l["telefone"], l["modelo"],
             l["tipo"], l["quantidade"], l["disponivel"], l["reservado"],
             _ids(l["ids"])]
            for l in linhas
        ),
    )

    resumo = {}
    for l in linhas:
        uf = resumo.setdefault(
            l["uf"], {"agentes": set(), "quantidade": 0, "disponivel": 0, "reservado": 0}
        )
        uf["agentes"].add(l["agente"])
        for campo in ("quantidade", "disponivel", "reservado"):
            uf[campo] += l[campo]
    _escrever(
        livro.create_sheet("Resumo por UF"),
        ["UF", "Agentes", "Quantidade", "Disponível", "Reservado"],
        (
            [uf, len(d["agentes"]), d["quantidade"], d["disponivel"], d["reservado"]]
            for uf, d in resumo.items()
        ),
    )

    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _ids(identificadores: list[str]) -> str:
    texto = ", ".join(identificadores)
    if len(texto) <= _LIMITE_CELULA:
        return texto
    corte = texto[:_LIMITE_CELULA].rsplit(", ", 1)[0]
    restantes = len(identificadores) - corte.count(", ") - 1
    return f"{corte} … (+{restantes})"


def _escrever(aba, cabecalho, linhas):
    aba.append(cabecalho)
    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = _CABECALHO
    for linha in linhas:
        aba.append(linha)
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    for indice, coluna in enumerate(aba.columns, start=1):
        largura = max(len(str(c.value or "")) for c in coluna)
        aba.column_dimensions[get_column_letter(indice)].width = min(max(largura + 2, 8), 60)
    if aba.max_column >= 10:
        for celula in aba["J"][1:]:
            celula.alignment = Alignment(wrap_text=True, vertical="top")
