"""Exportação de iscas com agentes, por região (UF)."""
from io import BytesIO

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from iscas.services import entrada as entrada_service
from iscas.services import solicitacao as solicitacao_service

pytestmark = pytest.mark.django_db


@pytest.fixture
def cenario(cliente, operador, agente, agente2, modelo_descartavel,
            modelo_retornavel, unidades_com_agente):
    agente2.uf = "ES"
    agente2.save(update_fields=["uf"])
    entrada_service.registrar_entrada(
        modelo=modelo_retornavel, identificadores=["R01", "R02"],
        destino=agente2, autor=operador,
    )
    pedido = solicitacao_service.abrir_solicitacao(
        cliente=cliente, itens=[(modelo_descartavel, 1)], autor=operador
    )
    solicitacao_service.criar_atribuicao(
        solicitacao=pedido, agente=agente,
        itens=[(modelo_descartavel, 1)], autor=operador,
    )


def _planilha(client, **params):
    resposta = client.get(reverse("iscas:saldo_excel"), params)
    assert resposta.status_code == 200
    return load_workbook(BytesIO(resposta.content))


def test_linhas_por_agente_e_modelo_ordenadas_por_uf(client, operador_logado, cenario):
    livro = _planilha(client)
    linhas = list(livro["Iscas com agentes"].iter_rows(min_row=2, values_only=True))

    es, sp = linhas
    assert es[:7] == ("ES", "São Paulo", "Agente Dois", "11999990000",
                      "Isca Retornável R2", "Retornável", 2)
    assert es[9] == "R01, R02"
    assert sp[0] == "SP" and sp[2] == "Agente Um"
    assert sp[6:9] == (8, 7, 1)
    assert sp[9] == ", ".join(f"A{i:03d}" for i in range(1, 9))

    resumo = list(livro["Resumo por UF"].iter_rows(min_row=2, values_only=True))
    assert resumo == [("ES", 1, 2, 2, 0), ("SP", 1, 8, 7, 1)]


def test_respeita_a_busca_da_tela(client, operador_logado, cenario):
    linhas = list(
        _planilha(client, q="dois")["Iscas com agentes"].iter_rows(min_row=2, values_only=True)
    )

    assert [l[2] for l in linhas] == ["Agente Dois"]
